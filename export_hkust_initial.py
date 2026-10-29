r"""
香港科技大学 (HKUST) 数据库切片与失效专业导出脚本
产出：
1. c:\Users\rtx\Desktop\香港科技大学_数据库切片全量清单(共77项).xlsx
2. c:\Users\rtx\Desktop\香港科技大学_数据库失效专业清单(is_valid=False_共21项).xlsx
3. c:\Users\rtx\Desktop\HKUST_步骤一至四\步骤1_数据库切片_db_slice.xlsx
4. c:\Users\rtx\Desktop\HKUST_步骤一至四\步骤1附表_数据库失效专业_is_valid_false.xlsx
"""

import os
import sys
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE_DIR, "data", "hkust")
os.makedirs(DATA_DIR, exist_ok=True)

DESKTOP_FOLDER = r"c:\Users\rtx\Desktop\HKUST_步骤一至四"
os.makedirs(DESKTOP_FOLDER, exist_ok=True)

slice_csv = os.path.join(DATA_DIR, "db_slice.csv")

# 样式
FONT_FAMILY = "微软雅黑"
HEADER_FILL = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
HEADER_FONT = Font(name=FONT_FAMILY, size=10, bold=True, color="FFFFFF")
INACTIVE_FILL = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")
GZ_FILL = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")

REGULAR_FONT = Font(name=FONT_FAMILY, size=9)
BOLD_FONT = Font(name=FONT_FAMILY, size=9, bold=True)
ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

THIN_BORDER = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

def build_excel(df, sheet_title):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.views.sheetView[0].showGridLines = True
    ws.row_dimensions[1].height = 28

    headers = [
        "序号", "数据库ID", "校区归属", "院校中文名", "院校英文名", "专业英文名",
        "专业中文名", "学位类型", "学制", "现有学费", "现有申请时间", "现有语言要求",
        "官网网址", "是否有效(is_valid)"
    ]
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = ALIGN_CENTER
        cell.border = THIN_BORDER

    for idx, r in df.iterrows():
        row_num = idx + 1
        db_id = r['id']
        prog_cn = str(r.get('program_name_cn', '') or '')
        prog_en = str(r.get('program_name_en', '') or '')
        url = str(r.get('program_website', '') or '')
        
        # 判断校区
        is_gz = ('hkust(gz)' in prog_en.lower() or 'hkust(gz)' in prog_cn.lower() or '广州' in prog_cn or 'hkust-gz' in url.lower())
        campus = "广州校区 (HKUST-GZ)" if is_gz else "清水湾本部 (Clearwater Bay)"

        is_valid = r.get('is_valid', True)
        is_valid_bool = (str(is_valid).lower() == 'true')

        vals = [
            row_num,
            db_id,
            campus,
            str(r.get('school_name_cn', '') or ''),
            str(r.get('school_name_en', '') or ''),
            prog_en,
            prog_cn,
            str(r.get('degree', '') or ''),
            str(r.get('program_duration', '') or ''),
            str(r.get('program_tuition', '') or ''),
            str(r.get('application_time', '') or '').replace('<br/>', '\n').strip(),
            str(r.get('language_requirements', '') or '').replace('<br/>', '\n').strip(),
            url,
            "有效 (True)" if is_valid_bool else "失效 (False)"
        ]
        ws.append(vals)

        cur_row = ws.max_row
        ws.row_dimensions[cur_row].height = 22
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=cur_row, column=col_idx)
            cell.font = REGULAR_FONT
            cell.border = THIN_BORDER
            if col_idx in [1, 2, 3, 8, 9, 14]:
                cell.alignment = ALIGN_CENTER
            else:
                cell.alignment = ALIGN_LEFT

            # 着色规则
            if not is_valid_bool:
                cell.fill = INACTIVE_FILL
            elif is_gz:
                cell.fill = GZ_FILL

    # 自动列宽
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or '')
            first_line = val_str.split('\n')[0]
            line_len = sum(2 if ord(c) > 127 else 1 for c in first_line)
            max_len = max(max_len, line_len)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 45)

    ws.auto_filter.ref = f"A1:N{ws.max_row}"
    return wb

def main():
    print("[*] 读取香港科技大学数据库切片 CSV...")
    df_all = pd.read_csv(slice_csv)
    df_false = df_all[df_all['is_valid'].astype(str).str.lower() == 'false'].reset_index(drop=True)
    df_true = df_all[df_all['is_valid'].astype(str).str.lower() == 'true'].reset_index(drop=True)

    print(f"    - 全量记录: {len(df_all)} 条")
    print(f"    - 有效记录 (is_valid=True): {len(df_true)} 条")
    print(f"    - 失效记录 (is_valid=False): {len(df_false)} 条")

    # 1. 全量表
    wb_all = build_excel(df_all, "HKUST_数据库全量切片(77项)")
    
    # 2. 失效表
    wb_false = build_excel(df_false, "HKUST_失效专业(is_valid=False)")

    # 导出文件
    out_all_desktop = r"c:\Users\rtx\Desktop\香港科技大学_数据库切片全量清单(共77项).xlsx"
    out_false_desktop = r"c:\Users\rtx\Desktop\香港科技大学_数据库失效专业清单(is_valid=False_共21项).xlsx"
    
    out_step1 = os.path.join(DESKTOP_FOLDER, "步骤1_数据库切片_db_slice.xlsx")
    out_step1_false = os.path.join(DESKTOP_FOLDER, "步骤1附表_数据库失效专业_is_valid_false.xlsx")

    out_data_all = os.path.join(DATA_DIR, "步骤1_数据库切片_db_slice.xlsx")
    out_data_false = os.path.join(DATA_DIR, "步骤1附表_数据库失效专业_is_valid_false.xlsx")

    wb_all.save(out_all_desktop)
    wb_false.save(out_false_desktop)
    wb_all.save(out_step1)
    wb_false.save(out_step1_false)
    wb_all.save(out_data_all)
    wb_false.save(out_data_false)

    print("\n[OK] 导出完成！已保存以下文件：")
    print(f"  1. 桌面全量切片: {out_all_desktop}")
    print(f"  2. 桌面失效清单: {out_false_desktop}")
    print(f"  3. 步骤一规范件: {out_step1}")
    print(f"  4. 步骤一附表件: {out_step1_false}")

if __name__ == "__main__":
    main()
