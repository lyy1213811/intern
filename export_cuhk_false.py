import os
import sys
import psycopg2
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

DB_CONFIG = {
    "host": "pgm-bp1k176i7sid0e775o.pg.rds.aliyuncs.com",
    "port": 5432,
    "user": "data_maintain",
    "password": "Str0ngP@5s!",
    "dbname": "tunshuedu_db"
}

REQUIRED_FIELDS = [
    "id",
    "school_name_cn",
    "school_name_en",
    "program_name_en",
    "program_name_cn",
    "program_code",
    "degree",
    "program_duration",
    "program_tuition",
    "application_time",
    "language_requirements",
    "program_website",
    "official_code",
    "is_valid"
]

def export_cuhk_false():
    print("[*] 正在连接数据库提取 [香港中文大学] is_valid = false 失效专业记录...")
    conn = psycopg2.connect(**DB_CONFIG)
    conn.set_session(readonly=True)

    fields_str = ", ".join(REQUIRED_FIELDS)
    sql = f"""
        SELECT {fields_str}
        FROM ai_selection_programs
        WHERE (school_name_cn = '香港中文大学' OR school_name_en = 'The Chinese University of Hong Kong')
          AND is_valid IS FALSE
        ORDER BY id ASC;
    """

    df = pd.read_sql(sql, conn)
    conn.close()

    print(f"[*] 成功查询到 {len(df)} 条失效专业记录。正在生成精美 Excel 文档...")

    # Excel 格式化
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CUHK_失效专业(is_valid=False)"

    # 表头
    headers = list(df.columns)
    ws.append(headers)

    # 数据行
    for row in df.itertuples(index=False):
        ws.append(list(row))

    # 样式配置
    font_family = "微软雅黑"
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name=font_family, size=10, bold=True, color="FFFFFF")
    regular_font = Font(name=font_family, size=9)
    badge_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid") # 浅红/橙色高亮
    badge_font = Font(name=font_family, size=9, bold=True, color="C00000")

    border_thin = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    ws.row_dimensions[1].height = 26
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center
        cell.border = border_thin

    # 启用自动筛选
    last_col_letter = get_column_letter(len(headers))
    ws.auto_filter.ref = f"A1:{last_col_letter}{ws.max_row}"

    is_valid_col_idx = headers.index("is_valid") + 1

    for row_idx in range(2, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 22
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = regular_font
            cell.border = border_thin

            val = cell.value
            if col_idx == 1: # id
                cell.alignment = align_center
            elif col_idx == is_valid_col_idx:
                cell.alignment = align_center
                cell.fill = badge_fill
                cell.font = badge_font
            elif isinstance(val, (int, float, bool)):
                cell.alignment = align_center
            elif str(val or '').startswith("http"):
                cell.alignment = align_left
            else:
                cell.alignment = align_left

    # 自动列宽调整
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or '')
            line_len = sum(2 if ord(c) > 127 else 1 for c in val_str.split('\n')[0])
            max_len = max(max_len, line_len)
        adjusted_width = min(max(max_len + 3, 10), 45)
        ws.column_dimensions[col_letter].width = adjusted_width

    # 输出路径
    desktop_file = r"c:\Users\rtx\Desktop\香港中文大学_数据库失效专业清单(is_valid=False_共16项).xlsx"
    wb.save(desktop_file)
    print(f"[OK] Excel 已成功导出到桌面: {desktop_file}")

    # 同时导出 CSV 备份
    data_dir = r"c:\Users\rtx\Desktop\intern\data\cuhk"
    os.makedirs(data_dir, exist_ok=True)
    csv_file = os.path.join(data_dir, "cuhk_db_slice_false.csv")
    df.to_csv(csv_file, index=False, encoding="utf-8-sig")
    print(f"[OK] CSV 备份已保存至: {csv_file}")

    return df

if __name__ == "__main__":
    export_cuhk_false()
