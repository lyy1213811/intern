r"""
香港理工大学 (PolyU) 步骤一：数据库切片全量与失效专业导出脚本
产出：
1. c:\Users\rtx\Desktop\香港理工大学_数据库切片全量清单(共153项).xlsx
2. c:\Users\rtx\Desktop\香港理工大学_数据库失效专业清单(is_valid=False_共25项).xlsx
3. c:\Users\rtx\Desktop\香港理工大学_数据库失效专业清单(is_valid=False_共25项).csv
4. c:\Users\rtx\Desktop\PolyU_步骤一至四\步骤1_数据库切片_db_slice.xlsx
5. c:\Users\rtx\Desktop\PolyU_步骤一至四\步骤1附表_数据库失效专业_is_valid_false.xlsx
"""

import os
import sys
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE_DIR, "data", "polyu")
os.makedirs(DATA_DIR, exist_ok=True)

DESKTOP_FOLDER = r"c:\Users\rtx\Desktop\PolyU_步骤一至四"
os.makedirs(DESKTOP_FOLDER, exist_ok=True)

slice_csv = os.path.join(DATA_DIR, "db_slice.csv")

# 样式
FONT_FAMILY = "微软雅黑"
HEADER_FILL = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
HEADER_FONT = Font(name=FONT_FAMILY, size=10, bold=True, color="FFFFFF")
INACTIVE_FILL = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")
WARN_FILL = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")

REGULAR_FONT = Font(name=FONT_FAMILY, size=9)
BOLD_FONT = Font(name=FONT_FAMILY, size=9, bold=True)
ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

THIN_BORDER = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

def build_excel(df, sheet_title):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.views.sheetView[0].showGridLines = True
    ws.row_dimensions[1].height = 28

    headers = [
        "序号", "数据库ID", "校区归属", "院校中文名", "院校英文名", "专业英文名",
        "专业中文名", "学位类型", "学制", "现有学费", "现有申请时间", "现有语言要求",
        "官网网址", "是否有效(is_valid)"
    ]
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = ALIGN_CENTER
        cell.border = THIN_BORDER

    for idx, r in df.iterrows():
        row_num = idx + 1
        db_id = r['id']
        prog_cn = str(r.get('program_name_cn', '') or '')
        prog_en = str(r.get('program_name_en', '') or '')
        url = str(r.get('program_website', '') or '')
        
        campus = "红磡本部 (Hung Hom)"
        is_valid = r.get('is_valid', True)
        is_valid_bool = (str(is_valid).lower() == 'true')

        vals = [
            row_num,
            db_id,
            campus,
            str(r.get('school_name_cn', '') or ''),
            str(r.get('school_name_en', '') or ''),
            prog_en,
            prog_cn,
            str(r.get('degree', '') or ''),
            str(r.get('program_duration', '') or ''),
            str(r.get('program_tuition', '') or ''),
            str(r.get('application_time', '') or '').replace('<br/>', '\n').strip(),
            str(r.get('language_requirements', '') or '').replace('<br/>', '\n').strip(),
            url,
            "有效 (True)" if is_valid_bool else "失效 (False)"
        ]
        ws.append(vals)

        cur_row = ws.max_row
        ws.row_dimensions[cur_row].height = 22

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=cur_row, column=col_idx)
            cell.font = REGULAR_FONT
            cell.border = THIN_BORDER
            if col_idx in [1, 2, 3, 8, 9, 14]:
                cell.alignment = ALIGN_CENTER
            else:
                cell.alignment = ALIGN_LEFT

            if not is_valid_bool:
                cell.fill = INACTIVE_FILL
                if col_idx == 14:
                    cell.font = BOLD_FONT

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or '')
            first_line = val_str.split('\n')[0]
            line_len = sum(2 if ord(c) > 127 else 1 for c in first_line)
            max_len = max(max_len, line_len)
        adjusted_width = min(max(max_len + 3, 10), 45)
        if col_letter in ['F', 'G']: adjusted_width = 32
        elif col_letter in ['K', 'L']: adjusted_width = 38
        ws.column_dimensions[col_letter].width = adjusted_width

    ws.auto_filter.ref = f"A1:N{ws.max_row}"
    return wb

def main():
    if not os.path.exists(slice_csv):
        print(f"[-] 文件不存在: {slice_csv}，请先运行 1_export_db_slice.py --school polyu")
        return

    print("[*] 读取 PolyU 数据库切片...")
    df = pd.read_csv(slice_csv)
    print(f"[*] 全量专业记录: {len(df)} 项")

    # 1. 导出全量切片
    wb_all = build_excel(df, "PolyU_数据库全量切片(153项)")
    out_all_desktop = r"c:\Users\rtx\Desktop\香港理工大学_数据库切片全量清单(共153项).xlsx"
    out_all_steps = os.path.join(DESKTOP_FOLDER, "步骤1_数据库切片_db_slice.xlsx")
    wb_all.save(out_all_desktop)
    wb_all.save(out_all_steps)
    print(f"[OK] 全量切片已生成至: {out_all_desktop}")
    print(f"[OK] 步骤一目录副本已保存至: {out_all_steps}")

    # 2. 导出失效专业清单 (is_valid == False)
    df_false = df[df['is_valid'].astype(str).str.lower() == 'false'].copy().reset_index(drop=True)
    print(f"[*] 失效专业记录: {len(df_false)} 项")

    wb_false = build_excel(df_false, "PolyU_失效专业(is_valid=False)")
    out_false_desktop = r"c:\Users\rtx\Desktop\香港理工大学_数据库失效专业清单(is_valid=False_共25项).xlsx"
    out_false_steps = os.path.join(DESKTOP_FOLDER, "步骤1附表_数据库失效专业_is_valid_false.xlsx")
    wb_false.save(out_false_desktop)
    wb_false.save(out_false_steps)
    print(f"[OK] 失效专业清单已生成至: {out_false_desktop}")
    print(f"[OK] 步骤一目录副本已保存至: {out_false_steps}")

    # 3. 导出失效专业 CSV 供快速核对
    out_false_csv = r"c:\Users\rtx\Desktop\香港理工大学_数据库失效专业清单(is_valid=False_共25项).csv"
    csv_rows = []
    for idx, r in df_false.iterrows():
        p_cn = str(r['program_name_cn'] or '')
        p_en = str(r['program_name_en'] or '')
        csv_rows.append({
            "序号": idx + 1,
            "数据库ID": r['id'],
            "校区": "红磡本部 (Hung Hom)",
            "院校中英文": f"{r['school_name_cn']} {r['school_name_en']}",
            "专业中英文 (专业名/学位)": f"{p_en} ({p_cn})",
            "学制": str(r.get('program_duration', '') or ''),
            "现有学费": str(r.get('program_tuition', '') or ''),
            "现有申请时间": str(r.get('application_time', '') or '').replace('<br/>', '\n').strip(),
            "现有语言要求": str(r.get('language_requirements', '') or '').replace('<br/>', '\n').strip(),
            "官网URL": str(r.get('program_website', '') or '')
        })
    pd.DataFrame(csv_rows).to_csv(out_false_csv, index=False, encoding='utf-8-sig')
    print(f"[OK] 失效专业 CSV 已生成至: {out_false_csv}")

if __name__ == "__main__":
    main()
