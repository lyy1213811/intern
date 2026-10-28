r"""
香港中文大学 (CUHK) 步骤五：最终交付件生成脚本
产出：c:\Users\rtx\Desktop\CUHK_数据核对与差异反馈表.xlsx
包含：
Sheet 1: 香港中文大学基准清单(149项TPg)
Sheet 2: 数据核对与差异反馈表 (覆盖：漏录、停招、其他异常[误标失效纠偏/轮次错位纠偏])
"""

import os
import re
import sys
import urllib.parse
from bs4 import BeautifulSoup
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE_DIR, "data", "cuhk")

db_slice_file = os.path.join(DATA_DIR, "db_slice.csv")
diff_file = os.path.join(DATA_DIR, "diff_programs.csv")
html_dir = os.path.join(DATA_DIR, "htmls")
out_desktop_file = r"c:\Users\rtx\Desktop\CUHK_数据核对与差异反馈表.xlsx"
out_data_file = os.path.join(DATA_DIR, "CUHK_数据核对与差异反馈表.xlsx")

# 样式定义
FONT_FAMILY = "微软雅黑"
HEADER_FILL = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
HEADER_FONT = Font(name=FONT_FAMILY, size=10, bold=True, color="FFFFFF")

STAGE1_FILL = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid") # 淡橙高亮漏录
STAGE2_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid") # 淡黄高亮其他异常/字段纠偏
SUSPEND_FILL = PatternFill(start_color="F2DCDB", end_color="F2DCDB", fill_type="solid") # 淡粉红高亮停招
INACTIVE_FILL = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")

REGULAR_FONT = Font(name=FONT_FAMILY, size=9)
BOLD_FONT = Font(name=FONT_FAMILY, size=9, bold=True)
ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

THIN_BORDER = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

def get_html_path(url):
    if not url or pd.isna(url): return None
    p = urllib.parse.urlparse(str(url))
    slug = p.fragment if p.fragment else os.path.basename(p.path.rstrip('/'))
    safe_slug = re.sub(r'[\\/*?:"<>|]', '_', slug).strip().lower()
    for f in os.listdir(html_dir):
        if not f.endswith('.html'): continue
        f_stem = f[:-5].lower()
        if f_stem == safe_slug or safe_slug.replace('-', '_') == f_stem.replace('-', '_'):
            return os.path.join(html_dir, f)
    return None

def parse_cuhk_html(fpath):
    if not fpath or not os.path.exists(fpath): return {}
    with open(fpath, 'r', encoding='utf-8', errors='ignore') as f:
        soup = BeautifulSoup(f.read(), 'html.parser')
    details = {}
    for t_box in soup.find_all('div', class_='title-bg'):
        title = t_box.get_text(strip=True)
        col4 = t_box.find_parent('div', class_='col-12')
        if col4:
            col8 = col4.find_next_sibling('div', class_='col-12')
            if col8:
                details[title] = col8.get_text('\n', strip=True)
    return details

def generate_cuhk_deliverables():
    print("[*] 开始生成香港中文大学 (CUHK) 最终质检成果表...")
    df_slice = pd.read_csv(db_slice_file)
    df_diff = pd.read_csv(diff_file)

    wb = openpyxl.Workbook()

    # ==========================================
    # Sheet 1: 香港中文大学基准清单(149项TPg)
    # ==========================================
    ws1 = wb.active
    ws1.title = "香港中文大学基准清单(149项TPg)"
    ws1.views.sheetView[0].showGridLines = True
    ws1.row_dimensions[1].height = 28

    headers1 = [
        "序号", "数据库ID", "所属学院", "院校中英文", "专业中英文 (专业名/学位)",
        "学制", "数据库现有学费", "数据库现有申请时间", "数据库现有语言要求", "官网URL"
    ]
    ws1.append(headers1)
    for col_idx in range(1, len(headers1) + 1):
        cell = ws1.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = ALIGN_CENTER
        cell.border = THIN_BORDER

    # 从 diff 表提取 faculty 映射
    fac_map = {}
    for _, r in df_diff.iterrows():
        if pd.notna(r.get('db_id')) and r.get('db_id') != '':
            try:
                fac_map[int(r['db_id'])] = str(r.get('faculty', '') or '')
            except:
                pass

    for idx, r in df_slice.iterrows():
        row_num = idx + 1
        db_id = r['id']
        school_str = f"{r['school_name_cn']} {r['school_name_en']}"
        prog_str = f"{r['program_name_en']} ({r['program_name_cn']})"
        faculty = fac_map.get(db_id, "研究生院/深圳校区" if "cuhk-shenzhen" in prog_str.lower() or "深圳" in prog_str else "研究生院/未分区")
        duration = str(r.get('program_duration', '') or '')
        tuition = str(r.get('program_tuition', '') or '')
        app_time = str(r.get('application_time', '') or '').replace('<br/>', '\n').strip()
        lang = str(r.get('language_requirements', '') or '').replace('<br/>', '\n').strip()
        url = str(r.get('program_website', '') or '')
        is_valid = r.get('is_valid', True)

        row_vals = [row_num, db_id, faculty, school_str, prog_str, duration, tuition, app_time, lang, url]
        ws1.append(row_vals)

        current_row = ws1.max_row
        ws1.row_dimensions[current_row].height = 22
        for col_idx in range(1, len(headers1) + 1):
            cell = ws1.cell(row=current_row, column=col_idx)
            cell.font = REGULAR_FONT
            cell.border = THIN_BORDER
            if col_idx in [1, 2, 6]:
                cell.alignment = ALIGN_CENTER
            elif col_idx in [3, 4, 5, 7, 8, 9, 10]:
                cell.alignment = ALIGN_LEFT

            if is_valid is False or str(is_valid).lower() == 'false':
                cell.fill = INACTIVE_FILL

    # Sheet 1 列宽
    for col in ws1.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or '')
            line_len = sum(2 if ord(c) > 127 else 1 for c in val_str.split('\n')[0])
            max_len = max(max_len, line_len)
        ws1.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 45)
    ws1.auto_filter.ref = f"A1:J{ws1.max_row}"

    # ==========================================
    # Sheet 2: 数据核对与差异反馈表
    # ==========================================
    ws2 = wb.create_sheet(title="数据核对与差异反馈表")
    ws2.views.sheetView[0].showGridLines = True
    ws2.row_dimensions[1].height = 28

    headers2 = [
        "序号", "院校中英文", "专业中英文", "官网URL", "数据库ID",
        "报错类型", "数据库现有值", "官网真实值", "官网原文依据", "备注", "审批意见"
    ]
    ws2.append(headers2)
    for col_idx in range(1, len(headers2) + 1):
        cell = ws2.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = ALIGN_CENTER
        cell.border = THIN_BORDER

    feedback_rows = []

    # ----------------------------------------------------
    # 类别一：全量漏录专业 (来自步骤3官网独有，共31项纯新增)
    # ----------------------------------------------------
    unmatched_off = df_diff[df_diff['category'].str.contains('官网独有', na=False)]
    
    # 历史被误标为失效的 4 项，稍后作为【其他异常-历史专业误标失效】单独录入
    reactivate_names = ['Juris Doctor', 'MBA', 'MA in Professional Educational Psychology', 'MSc in Clinical Gerontology and End-of-Life Care']
    
    for idx, r in unmatched_off.iterrows():
        off_name = r['official_name']
        off_url = r['official_url']
        faculty = r['faculty']
        
        if off_name in reactivate_names:
            continue

        feedback_rows.append({
            "stage_type": "漏录",
            "school": "香港中文大学 (The Chinese University of Hong Kong)",
            "prog_name": off_name,
            "url": off_url,
            "db_id": "无/新增",
            "err_type": "漏录",
            "db_val": "系统数据库中查无此在招课程",
            "off_val": f"2027/28 官方公布在招授课型硕士 (所属院系: {faculty})",
            "quote": f"CUHK Graduate School 官方招生列表: {off_name}",
            "note": "【阶段一确凿漏录】官网 2027/28 正式公布在招课程，系统数据库暂未收录，建议新增建档",
            "opinion": "待审批"
        })

    # ----------------------------------------------------
    # 类别二：停招/下线异常 (数据库在招但官网TPg无此课程，共2项)
    # ----------------------------------------------------
    suspend_programs = [
        {
            "id": 49672,
            "name": "Master of Fine Arts (艺术硕士)",
            "url": "https://www.gs.cuhk.edu.hk/programmes/arts/master-fine-arts",
            "db_val": "系统数据库中状态为在招 (is_valid=True)",
            "off_val": "官网授课型硕士(TPg)列表查无此专业",
            "quote": "CUHK Graduate School - Programmes - Arts (Master of Fine Arts 仅列于 Research Programmes)",
            "note": "【错录/已无授课型在招】官网授课型硕士仅开设艺术文学硕士 (MA in Fine Arts)；艺术硕士 MFA 属于研究型硕士 (RPg)，建议从授课型业务库中下线/停招"
        },
        {
            "id": 49673,
            "name": "Master of Music (音乐硕士)",
            "url": "https://www.gs.cuhk.edu.hk/programmes/arts/master-music",
            "db_val": "系统数据库中状态为在招 (is_valid=True)",
            "off_val": "官网授课型硕士(TPg)列表查无此专业",
            "quote": "CUHK Graduate School - Programmes - Arts (Master of Music 仅列于 Research Programmes)",
            "note": "【错录/已无授课型在招】官网授课型硕士仅开设音乐文学硕士 (MA in Music)；音乐硕士 MMus 属于研究型硕士 (RPg)，建议从授课型业务库中下线/停招"
        }
    ]
    for sp in suspend_programs:
        feedback_rows.append({
            "stage_type": "停招",
            "school": "香港中文大学 (The Chinese University of Hong Kong)",
            "prog_name": sp["name"],
            "url": sp["url"],
            "db_id": sp["id"],
            "err_type": "停招",
            "db_val": sp["db_val"],
            "off_val": sp["off_val"],
            "quote": sp["quote"],
            "note": sp["note"],
            "opinion": "待审批"
        })

    # ----------------------------------------------------
    # 类别三：其他异常 - 核心王牌专业历史记录误标失效 (共4项)
    # ----------------------------------------------------
    reactivate_items = [
        {
            "id": 2167,
            "name": "Juris Doctor (法律博士)",
            "url": "https://www.gs.cuhk.edu.hk/programmes/law/juris-doctor",
            "faculty": "Faculty of Law",
            "db_val": "历史ID: 2167 (库中状态被误标为 is_valid=False 失效)",
            "off_val": "2027/28 官方公布正式在招授课型专业 (院系: Faculty of Law)",
            "quote": "CUHK Graduate School - Programmes - Law: Juris Doctor",
            "note": "【核心在招专业误标失效】官网 2027/28 正常在招的核心学位课程，系统数据库历史记录误打上了失效标记导致前端不展示，建议解除失效状态重新激活或补录"
        },
        {
            "id": 2248,
            "name": "MBA (工商管理硕士)",
            "url": "https://www.gs.cuhk.edu.hk/programmes/business-administration/mba",
            "faculty": "CUHK Business School",
            "db_val": "历史ID: 2248 (库中状态被误标为 is_valid=False 失效)",
            "off_val": "2027/28 官方公布正式在招授课型专业 (院系: CUHK Business School)",
            "quote": "CUHK Graduate School - Programmes - Business: MBA",
            "note": "【核心在招专业误标失效】商学院核心旗舰 MBA 项目官网 2027/28 正常在招，库中历史记录被误标为失效，建议解除失效状态重新激活"
        },
        {
            "id": 2268,
            "name": "MA in Professional Educational Psychology (专业教育心理学文学硕士)",
            "url": "https://www.gs.cuhk.edu.hk/programmes/education/ma-professional-educational-psychology",
            "faculty": "Faculty of Education",
            "db_val": "历史ID: 2268 (库中状态被误标为 is_valid=False 失效)",
            "off_val": "2027/28 官方公布正式在招授课型专业 (院系: Faculty of Education)",
            "quote": "CUHK Graduate School - Programmes - Education: MA in Professional Educational Psychology",
            "note": "【在招专业误标失效】教育学院官网 2027/28 正在招生，数据库原记录被标记为失效，建议解除失效状态或重新建档"
        },
        {
            "id": 2278,
            "name": "MSc in Clinical Gerontology and End-of-Life Care (临床老年学与临终护理理学硕士)",
            "url": "https://www.gs.cuhk.edu.hk/programmes/medicine/msc-clinical-gerontology-and-end-life-care",
            "faculty": "Faculty of Medicine",
            "db_val": "历史ID: 2278 (库中状态被误标为 is_valid=False 失效)",
            "off_val": "2027/28 官方公布正式在招授课型专业 (院系: Faculty of Medicine)",
            "quote": "CUHK Graduate School - Programmes - Medicine: MSc in Clinical Gerontology and End-of-Life Care",
            "note": "【在招专业误标失效】医学院官网 2027/28 正常招生，库中历史记录被标记为失效导致前端漏报，建议解除失效状态重新激活"
        }
    ]
    for ri in reactivate_items:
        feedback_rows.append({
            "stage_type": "其他异常",
            "school": "香港中文大学 (The Chinese University of Hong Kong)",
            "prog_name": ri["name"],
            "url": ri["url"],
            "db_id": ri["id"],
            "err_type": "其他异常",
            "db_val": ri["db_val"],
            "off_val": ri["off_val"],
            "quote": ri["quote"],
            "note": ri["note"],
            "opinion": "待审批"
        })

    # ----------------------------------------------------
    # 类别四：其他异常 - 商学院提前批轮次错位与年份修正 (共4项)
    # ----------------------------------------------------
    early_round_items = [
        {
            "id": 2106,
            "name": "MSc in Actuarial Science and Insurance Analytics (精算学与保险分析理学硕士)",
            "url": "https://www.gs.cuhk.edu.hk/admissions/programme/business-administration#msc-in-actuarial-science-and-insurance-analytics",
            "db_val": "26年秋季入学: 开放申请(/) | Round 1 截止(2026-07-31) | Round 2 截止(2027-03-31)",
            "off_val": "2027/28 学年: 31 July 2026 (Early Round) | 31 March 2027 (Normal Round)",
            "quote": "Application Deadline: 31 July 2026 (Early Round); 31 March 2027",
            "note": "【入学年份与轮次命名错位纠偏】库中误将 2027/28 学年标为'26年秋季入学'；且官网 7月31日 明确为提前轮 (Early Round)，库中误将其标为 Round 1，导致所有轮次名称错位一轮"
        },
        {
            "id": 2128,
            "name": "MSc in Business Analytics (商业分析理学硕士)",
            "url": "https://www.gs.cuhk.edu.hk/admissions/programme/business-administration#msc-in-business-analytics",
            "db_val": "Round 1 截止(2026-07-31) | Round 2 截止(2027-03-31)",
            "off_val": "31 July 2026 (Early Round) | 31 March 2027 (Normal Round)",
            "quote": "Application Deadline: 31 July 2026 (Early Round); 31 March 2027",
            "note": "【轮次命名规范纠偏】官网明确设置 7月31日为提前轮(Early Round)，库中误将其标为 Round 1，建议规范轮次属性为提前批，防止误导申请人"
        },
        {
            "id": 2160,
            "name": "MSc in Marketing (市场营销理学硕士)",
            "url": "https://www.gs.cuhk.edu.hk/admissions/programme/business-administration#msc-in-marketing",
            "db_val": "Round 1 截止(2026-07-31) | Round 2 截止(2027-03-31)",
            "off_val": "31 July 2026 (Early Round) | 31 March 2027 (Normal Round)",
            "quote": "Application Deadline: 31 July 2026 (Early Round); 31 March 2027",
            "note": "【轮次命名规范纠偏】官网明确 7月31日为提前轮(Early Round)，库中误将其标为 Round 1，建议统一规范轮次命名"
        },
        {
            "id": 2140,
            "name": "Master of Accountancy (会计学硕士)",
            "url": "https://www.gs.cuhk.edu.hk/admissions/programme/business-administration#master-of-accountancy",
            "db_val": "Round 1 截止(2026-07-31) | Round 2 截止(2027-03-31)",
            "off_val": "31 July 2026 (Early Round) | 31 March 2027 (Normal Round)",
            "quote": "Application Deadline: 31 July 2026 (Early Round); 31 March 2027",
            "note": "【轮次命名规范纠偏】官网明确设置 7月31日为提前轮(Early Round)，库中误将其标为 Round 1，建议规范轮次名称"
        }
    ]
    for eri in early_round_items:
        feedback_rows.append({
            "stage_type": "其他异常",
            "school": "香港中文大学 (The Chinese University of Hong Kong)",
            "prog_name": eri["name"],
            "url": eri["url"],
            "db_id": eri["id"],
            "err_type": "其他异常",
            "db_val": eri["db_val"],
            "off_val": eri["off_val"],
            "quote": eri["quote"],
            "note": eri["note"],
            "opinion": "待审批"
        })

    # 填充 Sheet 2
    for seq, r_data in enumerate(feedback_rows, 1):
        row_vals = [
            seq, r_data['school'], r_data['prog_name'], r_data['url'], r_data['db_id'],
            r_data['err_type'], r_data['db_val'], r_data['off_val'], r_data['quote'],
            r_data['note'], r_data['opinion']
        ]
        ws2.append(row_vals)
        cur_row = ws2.max_row
        ws2.row_dimensions[cur_row].height = 24

        st = r_data['stage_type']
        if st == "漏录":
            fill = STAGE1_FILL
        elif st == "停招":
            fill = SUSPEND_FILL
        else:
            fill = STAGE2_FILL

        for c_idx in range(1, len(headers2) + 1):
            cell = ws2.cell(row=cur_row, column=c_idx)
            cell.font = REGULAR_FONT
            cell.border = THIN_BORDER
            cell.fill = fill
            if c_idx in [1, 5, 6, 11]:
                cell.alignment = ALIGN_CENTER
            else:
                cell.alignment = ALIGN_LEFT

    # Sheet 2 列宽与筛选
    for col in ws2.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or '')
            line_len = sum(2 if ord(c) > 127 else 1 for c in val_str.split('\n')[0])
            max_len = max(max_len, line_len)
        ws2.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 45)
    ws2.auto_filter.ref = f"A1:K{ws2.max_row}"

    # 保存
    wb.save(out_desktop_file)
    wb.save(out_data_file)
    print(f"\n[OK] 香港中文大学最终成果表导出完成！")
    print(f"    -> 桌面交付件: {out_desktop_file}")
    print(f"    -> 仓库备份件: {out_data_file}")
    print(f"    - Sheet 1 基准清单记录: {ws1.max_row - 1} 条")
    print(f"    - Sheet 2 质检反馈记录: {ws2.max_row - 1} 条")
    
    # 打印报错类型统计
    type_counts = pd.DataFrame(feedback_rows)['err_type'].value_counts()
    for t, c in type_counts.items():
        print(f"      * {t}: {c} 条")

if __name__ == "__main__":
    generate_cuhk_deliverables()
