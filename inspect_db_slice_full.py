# -*- coding: utf-8 -*-
import openpyxl, sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook(r'C:\Users\rtx\Desktop\HKBU_\步骤1_数据库切片_db_slice.xlsx', data_only=True)
ws = wb.active
headers = [ws.cell(1, c).value for c in range(1, ws.max_column+1)]

programs = []
for r in range(2, ws.max_row+1):
    prog = {headers[c-1]: ws.cell(r, c).value for c in range(1, ws.max_column+1)}
    programs.append(prog)

print(f'Total programs loaded from db_slice: {len(programs)}')
for p in programs:
    print(f"ID: {p['数据库ID']} | 校区: {p['办学归属']} | 有效: {p['是否有效(is_valid)']} | 中文: {p['专业中文名']}")
    print(f"    英文: {p['专业英文名']}")
    print(f"    学制: {p['学制']} | 学费: {p['现有学费']} | 时间: {p['现有申请时间']}")
    print(f"    URL: {p['官网网址']}")
    print("-" * 50)
