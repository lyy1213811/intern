r"""
香港浸会大学 (HKBU) 步骤一：数据库切片全量与失效专业导出脚本
产出：
1. c:\Users\rtx\Desktop\香港浸会大学_数据库切片全量清单(共62项).xlsx
2. c:\Users\rtx\Desktop\香港浸会大学_数据库失效专业清单(is_valid=False_共1项).xlsx
3. c:\Users\rtx\Desktop\香港浸会大学_数据库失效专业清单(is_valid=False_共1项).csv
4. c:\Users\rtx\Desktop\HKBU_步骤一至四\步骤1_数据库切片_db_slice.xlsx
5. c:\Users\rtx\Desktop\HKBU_步骤一至四\步骤1附表_数据库失效专业_is_valid_false.xlsx
6. intern\data\hkbu\db_slice.csv
"""

import os
import sys
import psycopg2
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = r"C:\Users\rtx\Desktop\intern"
DATA_DIR = os.path.join(BASE_DIR, "data", "hkbu")
os.makedirs(DATA_DIR, exist_ok=True)

STEPS_FOLDER = r"c:\Users\rtx\Desktop\HKBU_步骤一至四"
os.makedirs(STEPS_FOLDER, exist_ok=True)
DESKTOP_DIR = r"c:\Users\rtx\Desktop"

slice_csv = os.path.join(DATA_DIR, "db_slice.csv")

DB_CONFIG = {
    "host": os.getenv("DB_HOST", "pgm-bp1k176i7sid0e775o.pg.rds.aliyuncs.com"),
    "port": int(os.getenv("DB_PORT", 5432)),
    "user": os.getenv("DB_USER", "data_maintain"),
    "password": os.getenv("DB_PASSWORD", "Str0ngP@5s!"),
    "dbname": os.getenv("DB_NAME", "tunshuedu_db")
}

REQUIRED_FIELDS = [
    "id", "school_name_cn", "school_name_en", "program_name_en",
    "program_name_cn", "program_code", "degree", "program_duration",
    "program_tuition", "application_time", "language_requirements",
    "program_website", "official_code", "is_valid"
]

# 样式
FONT_FAMILY = "微软雅黑"
HEADER_FILL = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
HEADER_FONT = Font(name=FONT_FAMILY, size=10, bold=True, color="FFFFFF")
INACTIVE_FILL = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")
WARN_FILL = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
UIC_FILL = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")

REGULAR_FONT = Font(name=FONT_FAMILY, size=9)
BOLD_FONT = Font(name=FONT_FAMILY, size=9, bold=True)
ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

THIN_BORDER = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

def fetch_and_save_db():
    print("[*] 正在连接数据库导出 [香港浸会大学] 专业切片...")
    conn = psycopg2.connect(**DB_CONFIG)
    try:
        conn.set_session(readonly=True)
        fields_str = ", ".join(REQUIRED_FIELDS)
        sql = f"""
            SELECT {fields_str}
            FROM ai_selection_programs
            WHERE (school_name_cn = '香港浸会大学' OR school_name_en = 'Hong Kong Baptist University')
              AND degree = '硕士'
            ORDER BY id ASC
        """
        with conn.cursor() as cur:
            cur.execute(sql)
            rows = cur.fetchall()
            columns = [desc[0] for desc in cur.description]
            df = pd.DataFrame(rows, columns=columns)
            
        df.to_csv(slice_csv, index=False, encoding="utf-8-sig")
        print(f"[OK] 成功导出 {len(df)} 条专业记录至: {slice_csv}")
        return df
    finally:
        conn.close()

def build_excel(df, sheet_title):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.views.sheetView[0].showGridLines = True
    ws.row_dimensions[1].height = 28

    headers = [
        "序号", "数据库ID", "办学归属", "院校中文名", "院校英文名", "专业英文名",
        "专业中文名", "学位类型", "学制", "现有学费", "现有申请时间", "现有语言要求",
        "官网网址", "是否有效(is_valid)"
    ]
    ws.append(headers)
    for c_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = ALIGN_CENTER
        cell.border = THIN_BORDER

    for idx, r in df.reset_index(drop=True).iterrows():
        is_uic = 'uic.edu.cn' in str(r['program_website']).lower() or 'bnbu' in str(r['program_name_en']).lower() or '联合国际学院' in str(r['program_name_cn'])
        is_v = bool(r['is_valid'])
        campus = "UIC联合国际学院 (珠海)" if is_uic else "九龙塘本部"
        
        row_vals = [
            idx + 1, r['id'], campus, r['school_name_cn'], r['school_name_en'],
            r['program_name_en'], r['program_name_cn'], r['degree'],
            r['program_duration'], r['program_tuition'],
            str(r['application_time'] or '').replace('<br/>', '\n').strip(),
            str(r['language_requirements'] or '').replace('<br/>', '\n').strip(),
            r['program_website'], is_v
        ]
        ws.append(row_vals)
        c_row = ws.max_row
        ws.row_dimensions[c_row].height = 24
        for c_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=c_row, column=c_idx)
            cell.font = REGULAR_FONT
            cell.border = THIN_BORDER
            if c_idx in [1, 2, 3, 4, 8, 9, 14]:
                cell.alignment = ALIGN_CENTER
            else:
                cell.alignment = ALIGN_LEFT
            
            if not is_v:
                cell.fill = INACTIVE_FILL
            elif is_uic:
                cell.fill = UIC_FILL

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(sum(2 if ord(ch) > 127 else 1 for ch in str(cell.value or '').split('\n')[0]) for cell in col)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 11), 50)

    ws.auto_filter.ref = f"A1:N{ws.max_row}"
    return wb

def main():
    df = fetch_and_save_db()
    
    df_valid = df[df['is_valid'].astype(str).str.lower() == 'true'].copy()
    df_false = df[df['is_valid'].astype(str).str.lower() == 'false'].copy()
    
    print(f"[*] 统计分布: 总计 {len(df)} 项 | 有效 {len(df_valid)} 项 | 失效 {len(df_false)} 项")
    
    # 1. 导出全量切片 Excel (步骤1)
    wb_all = build_excel(df, "香港浸会大学数据库切片全量")
    wb_all.save(os.path.join(STEPS_FOLDER, "步骤1_数据库切片_db_slice.xlsx"))
    wb_all.save(os.path.join(DESKTOP_DIR, f"香港浸会大学_数据库切片全量清单(共{len(df)}项).xlsx"))
    print(f"[OK] 步骤1全量切片已生成至桌面及步骤目录！")
    
    # 2. 导出失效专业 Excel & CSV (步骤1附表)
    if len(df_false) > 0:
        wb_false = build_excel(df_false, "香港浸会大学数据库失效专业清单")
        wb_false.save(os.path.join(STEPS_FOLDER, "步骤1附表_数据库失效专业_is_valid_false.xlsx"))
        wb_false.save(os.path.join(DESKTOP_DIR, f"香港浸会大学_数据库失效专业清单(is_valid=False_共{len(df_false)}项).xlsx"))
        df_false.to_csv(os.path.join(DESKTOP_DIR, f"香港浸会大学_数据库失效专业清单(is_valid=False_共{len(df_false)}项).csv"), index=False, encoding="utf-8-sig")
        print(f"[OK] 步骤1附表失效专业已生成至桌面及步骤目录！")
    
if __name__ == "__main__":
    main()
