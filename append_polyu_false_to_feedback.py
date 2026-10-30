import os, sys, openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')

poly_dir = r"c:\Users\rtx\Desktop\PolyU_步骤一至四"
file_path = os.path.join(poly_dir, "香港理工大学_数据核对与差异反馈表.xlsx")
sub_file = os.path.join(poly_dir, "步骤1附表_数据库失效专业_is_valid_false.xlsx")

print(">>> Loading existing workbook:", file_path)
wb = openpyxl.load_workbook(file_path)
wb_sub = openpyxl.load_workbook(sub_file, data_only=True)
ws_sub = wb_sub.active

FONT_FAMILY = "微软雅黑"
HEADER_FILL = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid") # 深蓝
HEADER_FONT = Font(name=FONT_FAMILY, size=10, bold=True, color="FFFFFF")

# 行级底色定义
ADOPT_ROW_FILL = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid") # 浅绿色：采纳更新
MAINTAIN_ROW_FILL = PatternFill(start_color="F2DCDB", end_color="F2DCDB", fill_type="solid") # 浅粉色：驳回报错/确权维持

STATUS_ADOPT_FONT = Font(name=FONT_FAMILY, size=9, bold=True, color="276A3C") # 深绿字
STATUS_MAINTAIN_FONT = Font(name=FONT_FAMILY, size=9, bold=True, color="9C0006") # 深红字
REGULAR_FONT = Font(name=FONT_FAMILY, size=9)

THIN_BORDER = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

ws2 = wb['数据核对与差异反馈表']
existing_count = ws2.max_row - 1
print(f"Existing rows in Sheet 2: {existing_count}")

# 25 项历史失效条目的分类说明字典
false_explanations = {
    3541: "【历史双语班停招归档】商务分析中文辅以英文授课历史班次，现已停办归档，在招主专业由英文授课主条目承接。原库标为 is_valid=False 精确无误，确权维持历史归档免检。",
    3542: "【历史细分代码归并】商务分析历史细分招生代码，现已整合至主理学硕士课程架构。维持 is_valid=False 免检隔离。",
    3544: "【大类细分归档】经济学（应用经济学）细分方向，已归并至经济学大类招生。维持 is_valid=False 归档免检。",
    3545: "【大类细分归档】经济学（数字经济学）细分方向，已归并至经济学大类招生。维持 is_valid=False 归档免检。",
    3554: "【专业英文ESP子方向归并】专业英文文学硕士（英语教学子方向），已统一整合并入专业英文大类课程架构。维持 is_valid=False 免检归档。",
    3559: "【历史课程改制归并】高效能建筑理学硕士，属于工学院屋宇设备工程历史重组条目。维持 is_valid=False 归档免检。",
    3565: "【工学硕士改制归档】屋宇设备工程工学硕士 (MEng)，现已调整为理学硕士 (MSc) 轨制。维持 is_valid=False 归档免检。",
    3582: "【酒店学院SHTM组合课程归并】酒店业人工智能细分方向，SHTM 现已全系整合为国际酒店及旅游管理理学硕士组合课程架构。维持 is_valid=False 免检隔离。",
    3587: "【酒店学院SHTM组合课程归并】酒店创新与创业细分方向，已并入国际酒店及旅游管理组合课程。维持 is_valid=False 归档免检。",
    3588: "【酒店学院SHTM组合课程归并】国际酒店管理历史细分代码，已并入国际酒店及旅游管理组合课程。维持 is_valid=False 归档免检。",
    3589: "【酒店学院SHTM组合课程归并】国际旅游与活动管理细分方向，已并入国际酒店及旅游管理组合课程。维持 is_valid=False 归档免检。",
    3590: "【酒店学院SHTM组合课程归并】国际葡萄酒管理细分方向，已并入国际酒店及旅游管理组合课程。维持 is_valid=False 归档免检。",
    3591: "【酒店学院SHTM组合课程归并】奢华体验管理细分方向，已并入国际酒店及旅游管理组合课程。维持 is_valid=False 归档免检。",
    3595: "【工商管理历史单设代码归档】MBA 历史细分代码，现由主代码规范承接。维持 is_valid=False 归档免检。",
    3610: "【康复医疗专业隔年招生/改制归档】职业治疗学硕士 (MOT) 历史代码，现执行专业改版与隔年招生机制。维持 is_valid=False 归档免检。",
    3611: "【康复医疗专业隔年招生/改制归档】物理治疗学硕士 (MPT) 历史代码，现执行专业改版与隔年招生机制。维持 is_valid=False 归档免检。",
    3615: "【康复医疗专业改制归档】言语治疗硕士 (MST) 历史代码，现由新学年代码规范承接。维持 is_valid=False 归档免检。",
    3630: "【人工智能人文细分归并】生成式人工智能与人文科学（艺术与文化子方向），已整合并入大类理学硕士架构。维持 is_valid=False 免检隔离。",
    3634: "【专业英文ESP子方向归并】专业英文文学硕士（专业英语子方向），已并入主专业大类。维持 is_valid=False 归档免检。",
    3635: "【翻译硕士历史代码归并】翻译与传译文学硕士历史代码，现由在招主专业规范承接。维持 is_valid=False 归档免检。",
    3636: "【人工智能人文细分归并】生成式人工智能与人文科学（语言与传意子方向），已并入大类理学硕士架构。维持 is_valid=False 免检隔离。",
    3637: "【时装设计SFT子方向归并】时装设计文学硕士（时装实践子方向），SFT 学院现已整合为主专业大类招生。维持 is_valid=False 归档免检。",
    43565: "【时装设计SFT子方向归并】时装设计文学硕士（数码时装及创新子方向），已整合为主专业大类招生。维持 is_valid=False 归档免检。",
    44066: "【科技应用数学子方向归并】精算及投资科学细分方向，已统一整合并入应用数学主专业课程架构。维持 is_valid=False 归档免检。",
    49652: "【官方明确停办专业归档】人工智能与对外英语教学 (TESOL)，官方明确标注 Discontinued（停招停办）。原库标记 is_valid=False 100% 准确，坚决确权维持归档免检。"
}

# 遍历附表 25 项
added = 0
for r in range(2, ws_sub.max_row + 1):
    pid = ws_sub.cell(row=r, column=2).value # 数据库ID
    pname_en = ws_sub.cell(row=r, column=6).value
    pname_cn = ws_sub.cell(row=r, column=7).value
    url = ws_sub.cell(row=r, column=13).value
    
    prog_str = f"{pname_en} ({pname_cn})"
    school_str = "香港理工大学 The Hong Kong Polytechnic University"
    err_type = "历史失效专业免检"
    db_val = "is_valid=False (已归档)"
    sugg_val = "维持原值: is_valid=False (已合并归档/历史停招免检)"
    evidence = f"PolyU Official Postgraduate Catalogue / Programme Restructuring; Status: is_valid=False in DB."
    advise = false_explanations.get(pid, "【历史失效专业免检隔离】该专业属于历史停办或细分方向合并条目，原库标为 is_valid=False 符合红线5，予以免检隔离，确权维持历史归档。")
    
    existing_count += 1
    row_vals = [
        existing_count, school_str, prog_str, url, pid,
        err_type, db_val, sugg_val, evidence, advise
    ]
    ws2.append(row_vals)
    c_row = ws2.max_row
    ws2.row_dimensions[c_row].height = 24
    
    for c_idx in range(1, 11):
        cell = ws2.cell(row=c_row, column=c_idx)
        cell.fill = MAINTAIN_ROW_FILL # 浅粉色底色
        cell.font = REGULAR_FONT
        cell.border = THIN_BORDER
        if c_idx in [1, 5]: cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        elif c_idx == 6:
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.font = STATUS_MAINTAIN_FONT
        elif c_idx in [7, 8]: cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        else: cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    added += 1

print(f"Added {added} false records to Sheet 2. Total rows in Sheet 2 now: {ws2.max_row - 1}")

# 重新给整张 Sheet 2 设定样式与宽度
for col in ws2.columns:
    col_letter = get_column_letter(col[0].column)
    max_len = max(sum(2 if ord(ch) > 127 else 1 for ch in str(cell.value or '').split('\n')[0]) for cell in col)
    ws2.column_dimensions[col_letter].width = min(max(max_len + 3, 11), 50)

ws2.auto_filter.ref = f"A1:J{ws2.max_row}"

wb.save(file_path)
print(f"✅ Successfully updated: {file_path}")

# 同步保存至 Desktop 根目录备份
desktop_bkp = r"c:\Users\rtx\Desktop\香港理工大学_数据核对与差异反馈表.xlsx"
try:
    wb.save(desktop_bkp)
    print(f"✅ Synced to Desktop root: {desktop_bkp}")
except Exception as e:
    print(f"Notice: Desktop root file in use: {e}")
