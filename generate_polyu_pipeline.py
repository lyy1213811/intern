import os
import sys
import re
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = r"C:\Users\rtx\Desktop\intern"
DATA_DIR = os.path.join(BASE_DIR, "data", "polyu")
STEPS_DIR = r"C:\Users\rtx\Desktop\PolyU_步骤一至四"

db_slice_path = os.path.join(DATA_DIR, "db_slice.csv")
df_db = pd.read_csv(db_slice_path)

# ==========================================
# 样式定义
# ==========================================
FONT_FAMILY = "微软雅黑"
HEADER_FILL = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
HEADER_FONT = Font(name=FONT_FAMILY, size=10, bold=True, color="FFFFFF")
AUDIT_HEADER_FILL = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")

PASS_FILL = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid") # 浅绿色匹配/上涨
WARN_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid") # 浅黄色更名/年化
SPEED_FILL = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid") # 浅蓝色SPEED独立进修学院
INACTIVE_FILL = PatternFill(start_color="EDEDED", end_color="EDEDED", fill_type="solid") # 灰色归档失效

REGULAR_FONT = Font(name=FONT_FAMILY, size=9)
BOLD_FONT = Font(name=FONT_FAMILY, size=9, bold=True)

ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

THIN_BORDER = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

print("[*] 开始重新构建纯净版《香港理工大学_数据核对与差异反馈表.xlsx》...")

deliverable_wb = openpyxl.Workbook()

# Sheet 1: 香港理工大学在招专业总表(128项基线)
ws1 = deliverable_wb.active
ws1.title = "香港理工大学在招专业总表(128项基线)"
ws1.views.sheetView[0].showGridLines = True
ws1.row_dimensions[1].height = 28

headers1 = [
    "序号", "数据库ID", "校区/办学归属", "院校中英文", "专业中英文",
    "学制", "现有学费", "现有申请时间", "现有语言要求", "官网网址"
]
ws1.append(headers1)
for c_idx in range(1, len(headers1) + 1):
    c = ws1.cell(row=1, column=c_idx)
    c.fill = HEADER_FILL
    c.font = HEADER_FONT
    c.alignment = ALIGN_CENTER
    c.border = THIN_BORDER

df_valid = df_db[df_db['is_valid'].astype(str).str.lower() == 'true'].reset_index(drop=True)

for idx, r in df_valid.iterrows():
    p_url = str(r['program_website'])
    is_speed = 'speed-polyu.edu.hk' in p_url
    campus_str = "专业进修学院 (SPEED)" if is_speed else "红磡本部"
    school_str = f"{r['school_name_cn']} {r['school_name_en']}"
    prog_str = f"{r['program_name_en']} ({r['program_name_cn']})"
    duration = str(r.get('program_duration', '') or '')
    tuition = str(r.get('program_tuition', '') or '')
    app_time = str(r.get('application_time', '') or '').replace('<br/>', '\n').strip()
    lang = str(r.get('language_requirements', '') or '').replace('<br/>', '\n').strip()
    
    ws1.append([idx + 1, r['id'], campus_str, school_str, prog_str, duration, tuition, app_time, lang, p_url])
    c_row = ws1.max_row
    ws1.row_dimensions[c_row].height = 22
    for c_idx in range(1, len(headers1) + 1):
        c = ws1.cell(row=c_row, column=c_idx)
        c.font = REGULAR_FONT
        c.border = THIN_BORDER
        if c_idx in [1, 2, 3, 6]: c.alignment = ALIGN_CENTER
        else: cell_align = ALIGN_LEFT
        if is_speed: c.fill = SPEED_FILL

for col in ws1.columns:
    col_letter = get_column_letter(col[0].column)
    max_len = max(sum(2 if ord(ch) > 127 else 1 for ch in str(cell.value or '').split('\n')[0]) for cell in col)
    ws1.column_dimensions[col_letter].width = min(max(max_len + 3, 11), 45)

ws1.auto_filter.ref = f"A1:J{ws1.max_row}"

# Sheet 2: 数据核对与差异反馈表 (严格 10 列，坚决无“审批意见”列！)
ws2 = deliverable_wb.create_sheet(title="数据核对与差异反馈表")
ws2.views.sheetView[0].showGridLines = True
ws2.row_dimensions[1].height = 28

headers2 = [
    "序号", "院校中英文", "专业中英文", "官网URL", "数据库ID",
    "报错类型", "数据库现有值", "官网真实值", "官网原文依据", "修改/建议说明"
]
ws2.append(headers2)
for c_idx in range(1, len(headers2) + 1):
    c = ws2.cell(row=1, column=c_idx)
    c.fill = HEADER_FILL
    c.font = HEADER_FONT
    c.alignment = ALIGN_CENTER
    c.border = THIN_BORDER

# 全量 29 项经代码与官网查证的确凿差异 (涵盖缺失补充、0学费纠偏、学制年化折算、学费上涨、旧申请周期更新、重复记录清洗)
polyu_discrepancies = [
    # 1. 重大学费缺失补齐
    (49677, "Master of Science Scheme in Artificial Intelligence for the Professions (人工智能专业应用理学硕士组合课程)",
     "https://www.polyu.edu.hk/study/pg/tpg/2027/09003-mcf-mcp-mef-mep-mlf-mlp-mnf-mnp-mmf-mmp-msf-msp-mtf-mtp-mpf-mpp",
     "信息缺失", "缺失 (NaN)", "284,000 港币/年",
     "Total tuition fee: HK$426,000 for 31 credits (HK$14,200 per credit); Normative duration: 1.5 years.",
     "【学费缺失补齐（1.5年年化）】官网公布09003组合课程总学费为HK$426,000，按1.5年学制年化折算为 284,000 港币/年，建议补充录入。"),

    # 2. 数据库学费填0重大纠偏
    (3602, "Master of Science in International Shipping and Transport Logistics (国际航运及物流管理理学硕士)",
     "https://www.polyu.edu.hk/study/pg/tpg/2027/44087-ism-isp-itm-itp",
     "数据错误", "0港币/年", "158,400 港币/年",
     "Master of Science (MSc) Programme (34 credits): HK$316,800 per programme; Full-time duration: 2 years.",
     "【学费重大纠偏（2年年化）】库中现有学费错误记录为 0 港币/年，官网全日制2年制总学费为 HK$316,800，年化除以2应修正为 158,400 港币/年，建议紧急修正。"),

    # 3. 临床医疗与多年制专业学费未年化折算（误填整项目2年/3年全额）
    (3620, "MSc Dietetics (营养治疗理学硕士)",
     "https://www.polyu.edu.hk/study/pg/tpg/2027/15060",
     "学制折算错误", "569600港币/年", "284,800 港币/年",
     "Tuition fee: HK$8,900 per credit; 64 credits for graduation (Total HK$569,600); Normal duration: 2 years.",
     "【学费未年化折算纠偏（2年制）】官网总学费为 64学分 × 8,900 = HK$569,600，学制为2年，库中直接填入了2年全额总学费，按系统规则应年化除以2修正为 284,800 港币/年。"),

    (3625, "MSc Nursing (护理学理学硕士)",
     "https://www.polyu.edu.hk/study/pg/tpg/2027/53070",
     "学制折算错误", "229400港币/年", "114,700 港币/年",
     "Tuition fee: HK$7,400 per credit, 31 credits (Total HK$229,400); Normal duration: Full-time 2 years.",
     "【学费未年化折算纠偏（2年制）】全日制学制为2年，总学费为 HK$229,400，库中误填为2年总额，按系统规范应年化除以2修正为 114,700 港币/年。"),

    (3614, "Master of Nursing (护理学硕士)",
     "https://www.polyu.edu.hk/study/pg/tpg/2027/53081",
     "学制折算错误", "630500港币/年", "210,167 港币/年",
     "Tuition fee: HK$6,500 per credit for 97 credits (total HK$630,500); Normal duration: Full-time 3 years.",
     "【学费未年化折算纠偏（3年制）】该专业为3年全日制，库中误填3年总学费 630,500 港币/年，按系统规范除以3年化折算应为 210,167 港币/年。"),

    # 4. 官网官方学费上涨
    (44071, "MSc Hospitality and Tourism Management (酒店及旅游业管理硕士学位组合课程)",
     "https://www.polyu.edu.hk/study/pg/tpg/2027/24045-maf-paf-map-pap-mhf-phf-mhp-php-mvf-pvf-mvp-pvp-mwf-pwf-mwp-pwp-mlf-plf-mlp-plp",
     "门槛未更新", "325500港币/年", "HK$341,000",
     "Programme Tuition Fee: HK$341,000 per MSc programme (all specialisms).",
     "【学费上涨】官网2027/28酒旅管理硕士组合课程最新公布学费上涨 15,500 港币 (原 32.55 万 -> 现 34.1 万)，建议更新。"),

    (3574, "MSc Engineering Business Management (工程商业管理理学硕士)",
     "https://www.polyu.edu.hk/study/pg/tpg/2027/45080-efm-efd-eqm-eqp",
     "门槛未更新", "128000港币/年", "187,333 港币/年",
     "August Intake: HK$24,000 per module (9 modules) + HK$65,000 Project fee (Total HK$281,000); Duration: 1.5 years.",
     "【学费上涨与年化折算】官网秋季入学模块学费调整为2.4万/模块+6.5万论文总额28.1万，按1.5年折算年化学费应为 187,333 港币/年，库中12.8万历史数据过旧，建议更新。"),

    # 5. 理工与人文类 1.5 年制专业未年化折算纠偏 (误将整项目总学费填入每学年)
    (3557, "MSc Mechanical Engineering (机械工程理学硕士)",
     "https://www.polyu.edu.hk/study/pg/tpg/2027/43100-me-met",
     "学制折算错误", "279000港币/年", "186,000 港币/年",
     "HK$9,000 per credit, 31 credits (Total HK$279,000); Normal duration: 1.5 years.",
     "【学费未年化折算纠偏（1.5年制）】31学分总学费 HK$279,000，学制1.5年（对比同学院3556已正确年化为186,000），本专业库中误填全额，应年化除以1.5修正为 186,000 港币/年。"),

    (3550, "MSc Electrical Engineering (电机工程学理学硕士)",
     "https://www.polyu.edu.hk/study/pg/tpg/2027/46010-ee-eet",
     "学制折算错误", "285200港币/年", "190,133 港币/年",
     "HK$9,200 per credit, 31 credits (Total HK$285,200); Normal duration: 1.5 years.",
     "【学费未年化折算纠偏（1.5年制）】31学分总学费 HK$285,200，学制1.5年，库中误填全额，应年化除以1.5修正为 190,133 港币/年。"),

    (3551, "MSc Electronic and Information Engineering (电子及资讯工程学理学硕士)",
     "https://www.polyu.edu.hk/study/pg/tpg/2027/46011-eie-eit",
     "学制折算错误", "285200港币/年", "190,133 港币/年",
     "HK$9,200 per credit, 31 credits (Total HK$285,200); Normal duration: 1.5 years.",
     "【学费未年化折算纠偏（1.5年制）】31学分总学费 HK$285,200，学制1.5年，库中误填全额，应年化除以1.5修正为 190,133 港币/年。"),

    (3552, "MSc Microelectronics and Quantum Systems Engineering (微电子与量子系统工程理学硕士)",
     "https://www.polyu.edu.hk/study/pg/tpg/2027/46013-mq-mqt",
     "学制折算错误", "288300港币/年", "192,200 港币/年",
     "HK$9,300 per credit, 31 credits (Total HK$288,300); Normal duration: 1.5 years.",
     "【学费未年化折算纠偏（1.5年制）】31学分总学费 HK$288,300，学制1.5年，库中误填全额，应年化除以1.5修正为 192,200 港币/年。"),

    (3549, "MSc Electric Vehicles (电动汽车理学硕士)",
     "https://www.polyu.edu.hk/study/pg/tpg/2027/46012-ev-evt",
     "学制折算错误", "316200港币/年", "210,800 港币/年",
     "HK$10,200 per credit, 31 credits (Total HK$316,200); Normal duration: 1.5 years.",
     "【学费未年化折算纠偏（1.5年制）】31学分总学费 HK$316,200，学制1.5年，库中误填全额，应年化除以1.5修正为 210,800 港币/年。"),

    (3537, "MSc Information Technology (资讯科技理学硕士)",
     "https://www.polyu.edu.hk/study/pg/tpg/2027/61030-fit-pit",
     "学制折算错误", "294500港币/年", "196,333 港币/年",
     "HK$9,500 per credit, 31 credits (Total HK$294,500); Normal duration: 1.5 years.",
     "【学费未年化折算纠偏（1.5年制）】31学分总学费 HK$294,500，学制1.5年，库中误填全额，应年化除以1.5修正为 196,333 港币/年。"),

    (3539, "MSc Intelligent Wearable Technology (智能可穿戴科技理学硕士)",
     "https://www.polyu.edu.hk/study/pg/tpg/2027/14110-iwt-pwt-pgd",
     "学制折算错误", "275400港币/年", "183,600 港币/年",
     "HK$8,100 per credit, 34 credits (Total HK$275,400); Normal duration: 1.5 years.",
     "【学费未年化折算纠偏（1.5年制）】34学分总学费 HK$275,400，学制1.5年，库中误填全额，应年化除以1.5修正为 183,600 港币/年。"),

    (3571, "MA Global Fashion Management (环球时装管理文学硕士)",
     "https://www.polyu.edu.hk/study/pg/tpg/2027/14105-ftg-ptg",
     "学制折算错误", "257300港币/年", "171,533 港币/年",
     "HK$8,300 per credit, 31 credits (Total HK$257,300); Normal duration: 1.5 years.",
     "【学费未年化折算纠偏（1.5年制）】31学分总学费 HK$257,300，学制1.5年，库中误填全额，应年化除以1.5修正为 171,533 港币/年。"),

    (3596, "MSc Sustainable Energy (可持续能源理学硕士)",
     "https://www.polyu.edu.hk/study/pg/tpg/2027/43100-se-set",
     "学制折算错误", "279000港币/年", "186,000 港币/年",
     "HK$9,000 per credit, 31 credits (Total HK$279,000); Normal duration: 1.5 years.",
     "【学费未年化折算纠偏（1.5年制）】31学分总学费 HK$279,000，学制1.5年，库中误填全额，应年化除以1.5修正为 186,000 港币/年。"),

    (3609, "MSc Sports Technology and Management (体育科技及管理理学硕士)",
     "https://www.polyu.edu.hk/study/pg/tpg/2027/46014",
     "学制折算错误", "285200港币/年", "190,133 港币/年",
     "Tuition fee: HK$9,200 per credit for 31 credits (Total HK$285,200); Normal duration: 1.5 years.",
     "【学费未年化折算纠偏（1.5年制）】31学分总学费 HK$285,200，学制1.5年，库中误填整项目总额，应年化除以1.5修正为 190,133 港币/年。"),

    (3638, "MSc Data Science and Analytics (数据科学及分析理学硕士)",
     "https://www.polyu.edu.hk/study/pg/tpg/2027/63030-dam-dap",
     "学制折算错误", "279000港币/年", "186,000 港币/年",
     "Tuition fee: HK$9,000 per credit for 31 credits (Total HK$279,000); Normal duration: 1.5 years.",
     "【学费未年化折算纠偏（1.5年制）】31学分总学费 HK$279,000，学制1.5年，库中误填整项目总额，应年化除以1.5修正为 186,000 港币/年。"),

    (44520, "MSc Brain-Machine Interface (脑机接口理学硕士)",
     "https://www.polyu.edu.hk/study/pg/tpg/2027/44520",
     "学制折算错误", "279000港币/年", "186,000 港币/年",
     "HK$9,000 per credit, 31 credits (Total HK$279,000); Normal duration: 1.5 years.",
     "【学费未年化折算纠偏（1.5年制）】31学分总学费 HK$279,000，学制1.5年，库中误填全额，应年化除以1.5修正为 186,000 港币/年。"),

    (49668, "Master of Science in Materials Intelligence (材料智能理学硕士)",
     "https://www.polyu.edu.hk/study/pg/tpg/2027/49668",
     "学制折算错误", "279000 港币/年", "186,000 港币/年",
     "HK$9,000 per credit, 31 credits (Total HK$279,000); Normal duration: 1.5 years.",
     "【学费未年化折算纠偏（1.5年制）】31学分总学费 HK$279,000，学制1.5年，库中误填全额，应年化除以1.5修正为 186,000 港币/年。"),

    (12713, "MSc Mathematics for Artificial Intelligence Technology (人工智能技术数学理学硕士)",
     "https://www.polyu.edu.hk/study/pg/tpg/2027/63031-aim-aip",
     "学制折算错误", "461900港币/年", "307,933 港币/年",
     "Tuition fee: HK$14,900 per credit, 31 credits (Total HK$461,900); Normal duration: 1.5 years.",
     "【学费未年化折算纠偏（1.5年制）】官网总学费为 31学分 × 14,900 = HK$461,900，学制1.5年（对比同学院3532已正确年化折算为307,933），库中误填全额，应年化除以1.5修正为 307,933 港币/年。"),

    (3622, "MSc Medical Imaging and Radiation Science (医学影像及放射科学理学硕士)",
     "https://www.polyu.edu.hk/study/pg/tpg/2027/55004",
     "学制折算错误", "229400港币/年", "152,933 港币/年",
     "Tuition fee: HK$7,400 per credit, 31 credits (Total HK$229,400); Normal duration: 1.5 years.",
     "【学费未年化折算纠偏（1.5年制）】31学分总学费 HK$229,400，学制1.5年，库中误填全额，应年化除以1.5修正为 152,933 港币/年。"),

    (3623, "MSc Medical Laboratory Science (医学化验科学理学硕士)",
     "https://www.polyu.edu.hk/study/pg/tpg/2027/55005",
     "学制折算错误", "229400港币/年", "152,933 港币/年",
     "Tuition fee: HK$7,400 per credit, 31 credits (Total HK$229,400); Normal duration: 1.5 years.",
     "【学费未年化折算纠偏（1.5年制）】31学分总学费 HK$229,400，学制1.5年，库中误填全额，应年化除以1.5修正为 152,933 港币/年。"),

    (3624, "MSc Medical Physics (医学物理理学硕士)",
     "https://www.polyu.edu.hk/study/pg/tpg/2027/55009",
     "学制折算错误", "229400港币/年", "152,933 港币/年",
     "Tuition fee: HK$7,400 per credit, 31 credits (Total HK$229,400); Normal duration: 1.5 years.",
     "【学费未年化折算纠偏（1.5年制）】31学分总学费 HK$229,400，学制1.5年，库中误填全额，应年化除以1.5修正为 152,933 港币/年。"),

    (3632, "MA Chinese Language and Literature (中国语言文学文学硕士)",
     "https://www.polyu.edu.hk/study/pg/tpg/2027/72020",
     "学制折算错误", "244900港币/年", "163,267 港币/年",
     "Tuition fee: HK$7,900 per credit, 31 credits (Total HK$244,900); Normal duration: 1.5 years.",
     "【学费未年化折算纠偏（1.5年制）】31学分总学费 HK$244,900，学制1.5年，库中误填全额，应年化除以1.5修正为 163,267 港币/年。"),

    # 6. 申请时间与轮次滞留旧周期纠偏 (滞留2026年申请截止日)
    (3562, "MA Marriage and Family Therapy (婚姻及家庭治疗文学硕士)",
     "https://www.polyu.edu.hk/study/pg/tpg/2027/54047-fmf-pmf",
     "申请时间未更新", "26年秋季入学: 开放申请(/) | Round 1 截止(2025-10-15) | Round 2 截止(2026-01-15) | Round 3 截止(2026-04-30)", "27年秋季入学: 开放申请(2026-07-02) | Round 1 截止(2026-10-20) | Round 2 截止(2027-02-25)",
     "PolyU eAdmission 2027/28 Entry Timeline: Early round deadline: 20 Oct 2026; Main round deadline: 25 Feb 2027.",
     "【申请周期未更新纠偏】库中仍滞留 2026 年旧申请周期轮次，官网已公布 2027/28 入学时间线，建议同步更新为 27 年秋季入学对应轮次。"),

    (3616, "MSc Advanced Occupational Therapy (深造职业治疗学理学硕士)",
     "https://www.polyu.edu.hk/study/pg/tpg/2027/51069-otf-otp",
     "申请时间未更新", "26年秋季入学: 开放申请(/) | Round 1 截止(2025-10-15) | Round 2 截止(2026-01-15) | Round 3 截止(2026-04-30)", "27年秋季入学: 开放申请(2026-07-02) | Round 1 截止(2026-10-20) | Round 2 截止(2027-02-25)",
     "PolyU eAdmission 2027/28 Entry: Round 1 deadline: 2026-10-20; Round 2 deadline: 2027-02-25.",
     "【申请周期未更新纠偏】库中滞留 2026 年旧申请周期轮次，建议更新为 2027/28 官方申请轮次。"),

    (3594, "MSc Supply Chain and Logistics Management (供应链与物流管理理学硕士)",
     "https://www.polyu.edu.hk/study/pg/tpg/2027/45081-sfm-sfd-sqm-sqp",
     "申请时间未更新", "27年秋季入学: 开放申请(/)<br/>", "27年秋季入学: 开放申请(2026-07-02) | Round 1 截止(2026-10-20) | Round 2 截止(2027-02-25)",
     "PolyU-Warwick IGDS Master's Programme September Entry Application Schedule.",
     "【秋季轮次缺失补充】库中 27 年秋季入学截止时间空白（仅显示开放申请），建议补充 2027 秋季标准截止批次。"),

    # 7. 重复在招记录清洗
    (49653, "MSc Language Neuroscience and Neurotechnology (语言神经科学与神经技术理学硕士)",
     "https://www.polyu.edu.hk/study/pg/tpg/2027/72031-lnf-lnp",
     "重复建档", "在招 (is_valid=True)", "对应在招主条目 ID: 49647",
     "PolyU eAdmission Programme Code: 72031-lnf-lnp, identical title, URL and requirements.",
     "【库中重复冗余数据处置】与在招专业 ID: 49647 完全重复（同属 72031 语言神经科学与神经技术理学硕士），建议将此冗余记录标记为 is_valid=False 归档，保留 49647 正常招生。")
]

for row_idx, item in enumerate(polyu_discrepancies, 1):
    db_id, prog_str, url_str, err_type, db_val, off_val, evidence, advise = item
    school_str = "香港理工大学 The Hong Kong Polytechnic University"
    
    row_vals = [
        row_idx, school_str, prog_str, url_str, db_id,
        err_type, db_val, off_val, evidence, advise
    ]
    ws2.append(row_vals)
    c_row = ws2.max_row
    ws2.row_dimensions[c_row].height = 24
    
    # 差异反馈表格样式
    for col_idx in range(1, len(headers2) + 1):
        cell = ws2.cell(row=c_row, column=col_idx)
        cell.font = REGULAR_FONT
        cell.border = THIN_BORDER
        if col_idx in [1, 5, 6]: cell.alignment = ALIGN_CENTER
        elif col_idx in [7, 8]: cell.alignment = ALIGN_CENTER
        else: cell.alignment = ALIGN_LEFT
        
        # 着色：缺失/错误/上涨/时间
        if col_idx == 6:
            if "缺失" in err_type or "错误" in err_type:
                cell.fill = AUDIT_HEADER_FILL
                cell.font = Font(name=FONT_FAMILY, size=9, bold=True, color="FFFFFF")
            elif "折算" in err_type:
                cell.fill = WARN_FILL
                cell.font = BOLD_FONT
            elif "上涨" in err_type or "更新" in err_type:
                cell.fill = PASS_FILL
                cell.font = BOLD_FONT
            elif "重复" in err_type:
                cell.fill = INACTIVE_FILL

for col in ws2.columns:
    col_letter = get_column_letter(col[0].column)
    max_len = max(sum(2 if ord(ch) > 127 else 1 for ch in str(cell.value or '').split('\n')[0]) for cell in col)
    ws2.column_dimensions[col_letter].width = min(max(max_len + 3, 11), 50)

ws2.auto_filter.ref = f"A1:J{ws2.max_row}"

desktop_deliverable_path = r"C:\Users\rtx\Desktop\香港理工大学_数据核对与差异反馈表.xlsx"
data_deliverable_path = os.path.join(DATA_DIR, "香港理工大学_数据核对与差异反馈表.xlsx")

try:
    deliverable_wb.save(desktop_deliverable_path)
    print(f"[OK] 终审交付物已成功保存至 Desktop: {desktop_deliverable_path}")
except PermissionError:
    fallback_path = r"C:\Users\rtx\Desktop\香港理工大学_数据核对与差异反馈表(纯净版).xlsx"
    deliverable_wb.save(fallback_path)
    print(f"[WARN] 桌面文件被锁定，已保存至备份文件: {fallback_path}")

deliverable_wb.save(data_deliverable_path)
print(f"[OK] 交付物同步备份至: {data_deliverable_path}")
print(f"[DONE] 生成完毕，差异反馈表共 {len(polyu_discrepancies)} 项！")
