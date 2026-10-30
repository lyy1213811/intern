import os, sys, openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')

hkbu_dir = r"c:\Users\rtx\Desktop\HKBU_步骤一至四"
file_path = os.path.join(hkbu_dir, "香港浸会大学_数据核对与差异反馈表.xlsx")

wb = openpyxl.load_workbook(file_path)
ws2 = wb['数据核对与差异反馈表']

FONT_FAMILY = "微软雅黑"
HEADER_FILL = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid") # 深蓝
HEADER_FONT = Font(name=FONT_FAMILY, size=10, bold=True, color="FFFFFF")

# HKUST 标杆色系
TUITION_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid") # 淡黄：学费变动/年化纠偏
ROUND_FILL = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")   # 浅绿：申请时间/轮次补充
MAINTAIN_FILL = PatternFill(start_color="F2DCDB", end_color="F2DCDB", fill_type="solid")# 浅粉：驳回误判/确权维持

STATUS_TUITION_FONT = Font(name=FONT_FAMILY, size=9, bold=True, color="B25900") # 琥珀深棕
STATUS_ROUND_FONT = Font(name=FONT_FAMILY, size=9, bold=True, color="276A3C")   # 深绿
STATUS_MAINTAIN_FONT = Font(name=FONT_FAMILY, size=9, bold=True, color="9C0006")# 深红
REGULAR_FONT = Font(name=FONT_FAMILY, size=9)

THIN_BORDER = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

for r in range(2, ws2.max_row + 1):
    pid = ws2.cell(row=r, column=5).value
    err_type = str(ws2.cell(row=r, column=6).value or '')
    old_v = str(ws2.cell(row=r, column=7).value or '')
    sugg_v = str(ws2.cell(row=r, column=8).value or '')
    advise = str(ws2.cell(row=r, column=10).value or '')
    
    # 修复多年制建议值表达
    if pid in [8382, 43569, 43570, 49675] and "全项目总额" in sugg_v:
        sugg_v = f"维持原值: {old_v} (官网总额 ÷ 2年制)"
        ws2.cell(row=r, column=8, value=sugg_v)
        row_fill = MAINTAIN_FILL
        status_font = STATUS_MAINTAIN_FONT
    elif "年化" in err_type or "折算" in err_type or "学费" in err_type or "门槛" in err_type:
        if "纠偏" in err_type or "错误" in err_type or pid == 8377:
            row_fill = TUITION_FILL
            status_font = STATUS_TUITION_FONT
        else:
            row_fill = MAINTAIN_FILL
            status_font = STATUS_MAINTAIN_FONT
    elif "时间" in err_type or "信息缺失" in err_type or "轮次" in err_type:
        row_fill = ROUND_FILL
        status_font = STATUS_ROUND_FONT
    elif "免检" in err_type or "归档" in err_type or "维持" in advise:
        row_fill = MAINTAIN_FILL
        status_font = STATUS_MAINTAIN_FONT
    else:
        row_fill = ROUND_FILL
        status_font = STATUS_ROUND_FONT
        
    ws2.row_dimensions[r].height = 24
    for c_idx in range(1, 11):
        cell = ws2.cell(row=r, column=c_idx)
        cell.fill = row_fill
        cell.font = REGULAR_FONT
        cell.border = THIN_BORDER
        if c_idx in [1, 5]: cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        elif c_idx == 6:
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.font = status_font
        elif c_idx in [7, 8]: cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        else: cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

# 自动列宽
for col in ws2.columns:
    col_letter = get_column_letter(col[0].column)
    max_len = max(sum(2 if ord(ch) > 127 else 1 for ch in str(cell.value or '').split('\n')[0]) for cell in col)
    ws2.column_dimensions[col_letter].width = min(max(max_len + 3, 11), 50)

ws2.auto_filter.ref = f"A1:J{ws2.max_row}"

wb.save(file_path)
print("✅ HKUST-styled refinement complete for HKBU:", file_path)

desktop_bkp = r"c:\Users\rtx\Desktop\香港浸会大学_数据核对与差异反馈表.xlsx"
try:
    wb.save(desktop_bkp)
    print("✅ Synced to Desktop root:", desktop_bkp)
except Exception as e:
    print("Notice on desktop root:", e)
