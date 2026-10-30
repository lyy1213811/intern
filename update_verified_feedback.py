import os
import sys
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE_DIR, "data", "hkust")
slice_csv = os.path.join(DATA_DIR, "db_slice.csv")

# 样式
FONT_FAMILY = "微软雅黑"
HEADER_FILL = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
HEADER_FONT = Font(name=FONT_FAMILY, size=10, bold=True, color="FFFFFF")
STAGE2_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid") # 淡黄色学费上涨
ROUND_FILL = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")  # 浅绿色轮次更新
WARN_FILL = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")   # 浅橙色激活在招
GZ_FILL = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")

REGULAR_FONT = Font(name=FONT_FAMILY, size=9)
ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
THIN_BORDER = Border(
    left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9')
)

def update_verified_feedback_sheet():
    wb = openpyxl.load_workbook(r"c:\Users\rtx\Desktop\HKUST_数据核对与差异反馈表.xlsx")
    
    # 重新构建 Sheet 2: 数据核对与差异反馈表
    if "数据核对与差异反馈表" in wb.sheetnames:
        del wb["数据核对与差异反馈表"]
        
    ws2 = wb.create_sheet(title="数据核对与差异反馈表")
    ws2.views.sheetView[0].showGridLines = True
    ws2.row_dimensions[1].height = 28

    headers2 = [
        "序号", "院校中英文", "专业中英文", "官网URL", "数据库ID",
        "报错类型", "数据库现有值", "官网真实值", "官网原文依据", "修改/建议说明", "审批意见"
    ]
    ws2.append(headers2)
    for col_idx in range(1, len(headers2) + 1):
        cell = ws2.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = ALIGN_CENTER
        cell.border = THIN_BORDER

    # 33 项确凿学费上涨 + 2 项轮次追加 + 1 项休眠重启激活 = 共 36 条确凿差异项
    full_discrepancies = [
        # --- 商学院学费上涨 ---
        (3184, "Master of Science in Financial Technology (金融科技理学硕士)", "https://prog-crs.hkust.edu.hk/pgprog/2027-28/msc-fintech", "门槛未更新", "395000港币/年", "HK$405,000", "Program Fee: HK$405,000.", "【学费上涨】官网 2027/28 最新公布学费上涨 10,000 港币 (原 39.5 万 -> 现 40.5 万)，建议更新", "采纳 (Approved)"),
        (3185, "Master of Science in Finance (金融学理学硕士)", "https://prog-crs.hkust.edu.hk/pgprog/2027-28/msc-fina", "门槛未更新", "448800港币/年", "HK$460,000", "Program Fee: HK$460,000.", "【学费上涨】官网最新学费上调至 HK$460,000 (上涨 11,200 港币)，建议更新", "采纳 (Approved)"),
        (3195, "Master of Science in Accounting (会计学理学硕士)", "https://prog-crs.hkust.edu.hk/pgprog/2027-28/msc-acct", "门槛未更新", "390000港币/年", "HK$428,000", "Program Fee: HK$428,000.", "【学费大幅上涨】官方学费大幅上调 38,000 港币 (原 39 万 -> 现 42.8 万)，建议更新", "采纳 (Approved)"),
        (3192, "Master of Science in Business Analytics (商业分析理学硕士)", "https://prog-crs.hkust.edu.hk/pgprog/2027-28/msc-ba", "门槛未更新", "395000港币/年", "HK$415,000", "Program Fee: HK$415,000.", "【学费上涨】官方公布 2027/28 最新学费上涨 20,000 港币 (原 39.5 万 -> 现 41.5 万)，建议更新", "采纳 (Approved)"),
        (3193, "Master of Science in Economics (经济学理学硕士)", "https://prog-crs.hkust.edu.hk/pgprog/2027-28/msc-econ", "门槛未更新", "330000港币/年", "HK$350,000", "Program Fee: HK$350,000.", "【学费上涨】官方公布 2027/28 最新学费上涨 20,000 港币 (原 33 万 -> 现 35 万)，建议更新", "采纳 (Approved)"),
        (3209, "Master of Science in Marketing (市场营销理学硕士)", "https://prog-crs.hkust.edu.hk/pgprog/2027-28/msc-mark", "门槛未更新", "380000港币/年", "HK$400,000", "Program Fee: HK$400,000.", "【学费上涨】突破 40 万大关，上涨 20,000 港币 (原 38 万 -> 现 40 万)，建议更新", "采纳 (Approved)"),
        (3220, "Master of Science in Global Operations (全球运营理学硕士)", "https://prog-crs.hkust.edu.hk/pgprog/2027-28/msc-go", "门槛未更新", "340000港币/年", "HK$365,000", "Program Fee: HK$365,000.", "【学费上涨】官方最新学费上涨 25,000 港币 (原 34 万 -> 现 36.5 万)，建议更新", "采纳 (Approved)"),
        (3221, "Master of Science in International Management (国际管理理学硕士)", "https://prog-crs.hkust.edu.hk/pgprog/2027-28/msc-imt", "门槛未更新", "390000港币/年", "HK$420,000", "Program Fee: HK$420,000.", "【学费上涨】官方最新学费上涨 30,000 港币 (原 39 万 -> 现 42 万)，建议更新", "采纳 (Approved)"),
        (3217, "Master of Science in Information Systems Management (信息系统管理理学硕士)", "https://prog-crs.hkust.edu.hk/pgprog/2027-28/msc-ism", "门槛未更新", "336000港币/年", "Full-time: HK$360,000", "Tuition Fee: Full-time: HK$360,000; Part-time: HK$297,000.", "【学费上涨】商学院信息系统管理全日制学费上调 24,000 港币 (原 33.6 万 -> 现 36 万)，建议更新", "采纳 (Approved)"),
        (3219, "Master of Science in Family Office and Family Business (家族办公室与家族企业理学硕士)", "https://prog-crs.hkust.edu.hk/pgprog/2027-28/msc-fofb", "门槛未更新", "463500港币/年", "HK$477,400", "Program Fee: HK$477,400.", "【学费上涨】家族办公室硕士学费上调 13,900 港币 (原 46.35 万 -> 现 47.74 万)，建议更新", "采纳 (Approved)"),

        # --- 工学院学费上涨 ---
        (3191, "Master of Science in Information Technology (信息技术理学硕士)", "https://prog-crs.hkust.edu.hk/pgprog/2027-28/msc-it", "门槛未更新", "300000港币/年", "Full-time: HK$320,000", "Tuition Fee: Full-time: HK$320,000.", "【学费上涨】计算机系信息技术全日制硕士学费上涨 20,000 港币 (原 30 万 -> 现 32 万)，建议更新", "采纳 (Approved)"),
        (3196, "Master of Science in Electronic Engineering (电子工程理学硕士)", "https://prog-crs.hkust.edu.hk/pgprog/2027-28/msc-eleg", "门槛未更新", "240000港币/年", "Full-time: HK$298,000", "Tuition Fee: Full-time: HK$298,000.", "【学费大幅上涨】电子工程全日制学费大幅上调 58,000 港币 (+24.2%，原 24 万 -> 现 29.8 万)，建议更新", "采纳 (Approved)"),
        (3197, "Master of Science in Integrated Circuits (集成电路理学硕士)", "https://prog-crs.hkust.edu.hk/pgprog/2027-28/msc-ic", "门槛未更新", "240000港币/年", "Full-time: HK$298,000", "Tuition Fee: Full-time: HK$298,000.", "【学费大幅上涨】芯片方向全日制学费大幅上调 58,000 港币 (原 24 万 -> 现 29.8 万)，建议更新", "采纳 (Approved)"),
        (3198, "Master of Science in Telecommunications (电信学理学硕士)", "https://prog-crs.hkust.edu.hk/pgprog/2027-28/msc-telecom", "门槛未更新", "240000港币/年", "Full-time: HK$298,000", "Tuition Fee: Full-time: HK$298,000.", "【学费大幅上涨】电信学全日制硕士学费大幅上调 58,000 港币 (原 24 万 -> 现 29.8 万)，建议更新", "采纳 (Approved)"),
        (3246, "Master of Science in Big Data Technology (大数据技术理学硕士)", "https://prog-crs.hkust.edu.hk/pgprog/2027-28/msc-bdt", "门槛未更新", "330000港币/年", "Full-time: HK$350,000", "Tuition Fee: Full-time: HK$350,000.", "【学费上涨】大数据技术全日制硕士学费上涨 20,000 港币 (原 33 万 -> 现 35 万)，建议更新", "采纳 (Approved)"),
        (3211, "Master of Science in Chemical and Energy Engineering (化学与能源工程理学硕士)", "https://prog-crs.hkust.edu.hk/pgprog/2027-28/msc-cee", "门槛未更新", "208000港币/年", "Full-time: HK$229,000", "Tuition Fee: Full-time: HK$229,000.", "【学费上涨】化学与能源全日制硕士学费上涨 21,000 港币 (原 20.8 万 -> 现 22.9 万)，建议更新", "采纳 (Approved)"),
        (3247, "Master of Science in Biomolecular Engineering (生物分子工程与健康信息学硕士)", "https://prog-crs.hkust.edu.hk/pgprog/2027-28/msc-behi", "门槛未更新", "208000港币/年", "Full-time: HK$229,000", "Tuition Fee: Full-time: HK$229,000.", "【学费上涨】生物分子工程硕士学费上涨 21,000 港币 (原 20.8 万 -> 现 22.9 万)，建议更新", "采纳 (Approved)"),
        (3204, "Master of Science in Mechanical Engineering (机械工程理学硕士)", "https://prog-crs.hkust.edu.hk/pgprog/2027-28/msc-mech", "门槛未更新", "216500港币/年", "Full-time: HK$249,000", "Tuition Fee: Full-time: HK$249,000; Part-time: HK$216,000.", "【学费上涨】机械工程全日制学费上调 32,500 港币 (原 21.65 万 -> 现 24.9 万)，建议更新", "采纳 (Approved)"),
        (3208, "Master of Science in Intelligent Building Technology and Management (智能建筑物技术与管理理学硕士)", "https://prog-crs.hkust.edu.hk/pgprog/2027-28/msc-ibtm", "门槛未更新", "216500港币/年", "Full-time: HK$249,000", "Tuition Fee: Full-time: HK$249,000; Part-time: HK$216,000.", "【学费上涨】智能建筑全日制学费上调 32,500 港币 (原 21.65 万 -> 现 24.9 万)，建议更新", "采纳 (Approved)"),
        (3235, "Master of Science in Aeronautical Engineering (航空工程理学硕士)", "https://prog-crs.hkust.edu.hk/pgprog/2027-28/msc-ae", "门槛未更新", "216500港币/年", "Full-time: HK$249,000", "Tuition Fee: Full-time: HK$249,000; Part-time: HK$216,000.", "【学费上涨】航空工程全日制学费上调 32,500 港币 (原 21.65 万 -> 现 24.9 万)，建议更新", "采纳 (Approved)"),
        (3232, "Master of Science in Technology Leadership and Entrepreneurship (科技领导与创业理学硕士)", "https://prog-crs.hkust.edu.hk/pgprog/2027-28/msc-tle", "门槛未更新", "146667港币/年", "HK$300,000 (1.5年总额，折合 200,000 港币/年)", "Program Fee: HK$300,000 (Normative Duration: 1.5 years).", "【学费上涨（1.5年年化）】官网2027/28总学费由22万上调至30万，按1.5年学制折算年化学费为 200,000 港币/年 (上涨 53,333 港币/年)，建议更新", "采纳 (Approved)"),

        # --- 理学院学费上涨 ---
        (3183, "Master of Science in Financial Mathematics (金融数学理学硕士)", "https://prog-crs.hkust.edu.hk/pgprog/2027-28/msc-mafm", "门槛未更新", "166667港币/年", "HK$300,000 (1.5年总额，折合 200,000 港币/年)", "Program Fee: HK$300,000 (Normative Duration: 1.5 years).", "【学费上涨（1.5年年化）】官网2027/28总学费由25万上调至30万，按1.5年学制折算年化学费为 200,000 港币/年 (上涨 33,333 港币/年)，建议更新", "采纳 (Approved)"),
        (3203, "Master of Science in Mathematics for Educators (教育数学理学硕士)", "https://prog-crs.hkust.edu.hk/pgprog/2027-28/msc-maed", "门槛未更新", "150000港币/年", "HK$175,000", "Program Fee: HK$175,000.", "【学费上涨】数学系教育数学学费上调 25,000 港币 (原 15 万 -> 现 17.5 万)，建议更新", "采纳 (Approved)"),
        (3248, "Master of Science in Biotechnology (生物技术理学硕士)", "https://prog-crs.hkust.edu.hk/pgprog/2027-28/msc-biot", "门槛未更新", "142000港币/年", "Full-time: HK$168,000", "Tuition Fee: Full-time: HK$168,000.", "【学费上涨】生命科学系生物技术硕士学费上涨 26,000 港币 (原 14.2 万 -> 现 16.8 万)，建议更新", "采纳 (Approved)"),
        (3245, "Master of Science in Analytical Chemistry (分析化学理学硕士)", "https://prog-crs.hkust.edu.hk/pgprog/2027-28/msc-anchem", "门槛未更新", "160000港币/年", "HK$170,000", "Tuition Fee: HK$170,000.", "【学费上涨】化学系分析化学硕士官方最新学费上调 10,000 港币 (原 16 万 -> 现 17 万)，建议更新", "采纳 (Approved)"),
        (3249, "Master of Science in Data-Driven Modeling (数据驱动建模理学硕士)", "https://prog-crs.hkust.edu.hk/pgprog/2027-28/msc-ddm", "门槛未更新", "260000港币/年", "HK$286,000", "Program Fee: $286,000.", "【学费上涨】物理与数学交叉数据建模硕士学费上调 26,000 港币 (原 26 万 -> 现 28.6 万)，建议更新", "采纳 (Approved)"),
        (3225, "Master of Science in Environmental Health and Safety (环境健康与安全理学硕士)", "https://prog-crs.hkust.edu.hk/pgprog/2027-28/msc-ehs", "门槛未更新", "165000港币/年", "HK$175,000", "Program Fee: HK$175,000.", "【学费上涨】环境健康与安全硕士学费上调 10,000 港币 (原 16.5 万 -> 现 17.5 万)，建议更新", "采纳 (Approved)"),
        (3226, "Master of Science in Environmental Science, Management and Sustainability (环境科学、管理与可持续发展理学硕士)", "https://prog-crs.hkust.edu.hk/pgprog/2027-28/msc-evsm", "门槛未更新", "120000港币/年", "HK$198,000", "Program Fee: HK$198,000.", "【学费大幅上涨】官方最新公布学费大幅上调 78,000 港币 (原 12 万 -> 现 19.8 万)，建议更新", "采纳 (Approved)"),
        (3239, "Master of Science in Drug Regulatory Affairs and Policy (药物监管事务与政策理学硕士)", "https://prog-crs.hkust.edu.hk/pgprog/2027-28/msc-drap", "门槛未更新", "300000港币/年", "HK$330,000", "Program Fee: HK$330,000.", "【学费上涨】药物监管硕士学费上调 30,000 港币 (原 30 万 -> 现 33 万)，建议更新", "采纳 (Approved)"),

        # --- 人文社科学院学费上涨 ---
        (3240, "Master of Arts in Chinese Culture (中国文化文学硕士)", "https://prog-crs.hkust.edu.hk/pgprog/2027-28/ma-cc", "门槛未更新", "145000港币/年", "HK$174,000", "Tuition Fee: HK$174,000.", "【学费上涨】人文社科学院中国文化硕士学费上涨 29,000 港币 (原 14.5 万 -> 现 17.4 万)，建议更新", "采纳 (Approved)"),
        (44059, "Master of Arts in International Language Education (国际语言教育文学硕士)", "https://prog-crs.hkust.edu.hk/pgprog/2027-28/ma-ile", "门槛未更新", "135000港币/年", "HK$165,000", "Tuition Fee: HK$165,000.", "【学费上涨】国际语言教育硕士官方公布学费上调 30,000 港币 (原 13.5 万 -> 现 16.5 万)，建议更新", "采纳 (Approved)"),
        (3212, "Master of Arts in Social Science (社会学文学硕士)", "https://prog-crs.hkust.edu.hk/pgprog/2027-28/pgd-ma-sosc", "门槛未更新", "175000港币/年", "HK$178,000", "Program Fee: MA: HK$178,000.", "【学费上涨】社会学文学硕士全日制学费上调 3,000 港币 (原 17.5 万 -> 现 17.8 万)，建议更新", "采纳 (Approved)"),
        (3215, "Master of Science in Global China Studies (全球中国研究理学硕士)", "https://prog-crs.hkust.edu.hk/pgprog/2027-28/msc-gcs", "门槛未更新", "160000港币/年", "HK$170,000", "Program Fee: HK$170,000.", "【学费上涨】全球中国研究硕士学费上调 10,000 港币 (原 16 万 -> 现 17 万)，建议更新", "采纳 (Approved)"),

        # --- 申请轮次未同步更新 ---
        (3247, "Master of Science in Biomolecular Engineering (生物分子工程与健康信息学硕士)", "https://prog-crs.hkust.edu.hk/pgprog/2027-28/msc-behi", "申请时间未更新", "27年秋季: 仅Round 1(2026-12-01)", "追加 Round 2: 2027-02-15", "Application Deadlines: Non-local Applicants Full-time: 1 Dec 2026 (Round 1); 15 Feb 2027 (Round 2)", "【轮次追加】官网2027/28秋季全日制已明确设立 Round 2 (2027-02-15)，库中仅录入 Round 1，建议补齐轮次", "采纳 (Approved)"),
        (3235, "Master of Science in Aeronautical Engineering (航空工程理学硕士)", "https://prog-crs.hkust.edu.hk/pgprog/2027-28/msc-ae", "申请时间未更新", "27年秋季: 录至 Round 2(2027-03-01)", "追加 Round 3: 2027-06-01", "Application Deadlines: Non-local Applicants Full-time: 1 Dec 2026 (Round 1); 1 Mar 2027 (Round 2); 1 Jun 2027 (Round 3)", "【轮次追加】官网2027/28秋季全日制已明确设立 Round 3 (2027-06-01)，库中缺失第3轮，建议补齐", "采纳 (Approved)"),

        # --- 休眠重启激活 ---
        (44058, "Master of Science in Global Marine Resources Management (环球海洋资源管理理学硕士)", "https://prog-crs.hkust.edu.hk/pgprog/2027-28/msc-gmrm", "休眠重启激活", "历史ID: 44058 (库中状态为 is_valid=False 失效)", "2027/28 官方公布正式在招全日制双学位 (与英国南安普顿大学联合开设，Full-time 1年)", "Master of Science in Global Marine Resources Management (Dual degrees by HKUST and University of Southampton). Full-time 1 year. Instalment to HKUST: HK$140,000; Southampton: £10,000.", "【休眠重启专业建议激活】该专业因历史年度暂停招生被归档为 is_valid=False，2027/28 官方招生目录正式重启全日制招生，建议解除失效状态重新激活入库！（注：库中另有ID 3238为同项目，建议激活44058并保留3238作为归档去重）", "采纳 (Approved - 激活在招)")
    ]

    for seq, item in enumerate(full_discrepancies, 1):
        pid, pname, purl, etype, dval, oval, quote, note, opinion = item
        row_vals = [
            seq,
            "香港科技大学 (The Hong Kong University of Science and Technology)",
            pname,
            purl,
            pid,
            etype,
            dval,
            oval,
            quote,
            note,
            opinion
        ]
        ws2.append(row_vals)
        cur_row = ws2.max_row
        ws2.row_dimensions[cur_row].height = 24

        if "激活" in etype:
            fill = WARN_FILL
        elif "申请时间" in etype:
            fill = ROUND_FILL
        else:
            fill = STAGE2_FILL

        for c_idx in range(1, len(headers2) + 1):
            cell = ws2.cell(row=cur_row, column=c_idx)
            cell.font = REGULAR_FONT
            cell.border = THIN_BORDER
            cell.fill = fill
            if c_idx in [1, 5, 6, 11]: cell.alignment = ALIGN_CENTER
            else: cell.alignment = ALIGN_LEFT

    for col in ws2.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or '')
            first_line = val_str.split('\n')[0]
            line_len = sum(2 if ord(c) > 127 else 1 for c in first_line)
            max_len = max(max_len, line_len)
        ws2.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 45)
    ws2.auto_filter.ref = f"A1:K{ws2.max_row}"

    out_path = r"c:\Users\rtx\Desktop\HKUST_数据核对与差异反馈表.xlsx"
    wb.save(out_path)
    wb.save(os.path.join(DATA_DIR, "HKUST_数据核对与差异反馈表.xlsx"))
    print(f"[OK] 已将 36 项全面精准质检项更新至: {out_path}")
    print(f"    - 33 项确凿学费上涨 (含1.5年年化折算)")
    print(f"    - 2 项申请轮次追加")
    print(f"    - 1 项休眠重启激活 (GMRM)")

if __name__ == "__main__":
    update_verified_feedback_sheet()
