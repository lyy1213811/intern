import os, sys, openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

sys.stdout.reconfigure(encoding='utf-8')

poly_dir = r"c:\Users\rtx\Desktop\PolyU_步骤一至四"
file_path = os.path.join(poly_dir, "香港理工大学_数据核对与差异反馈表.xlsx")

wb = openpyxl.load_workbook(file_path)
ws2 = wb['数据核对与差异反馈表']

FONT_FAMILY = "微软雅黑"
HEADER_FILL = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid") # 深蓝
HEADER_FONT = Font(name=FONT_FAMILY, size=10, bold=True, color="FFFFFF")

ADOPT_ROW_FILL = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid") # 浅绿色：采纳更新
MAINTAIN_ROW_FILL = PatternFill(start_color="F2DCDB", end_color="F2DCDB", fill_type="solid") # 浅粉色：驳回报错/确权维持

STATUS_ADOPT_FONT = Font(name=FONT_FAMILY, size=9, bold=True, color="276A3C") # 深绿字
STATUS_MAINTAIN_FONT = Font(name=FONT_FAMILY, size=9, bold=True, color="9C0006") # 深红字
REGULAR_FONT = Font(name=FONT_FAMILY, size=9)

THIN_BORDER = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

for r in range(2, ws2.max_row + 1):
    err_type = str(ws2.cell(row=r, column=6).value or '')
    advise = str(ws2.cell(row=r, column=10).value or '')
    
    # 判定是采纳更新还是确权维持/驳回
    if "免检" in err_type or "核验" in err_type or "重复" in err_type or "维持" in advise:
        is_adopt = False
    else:
        is_adopt = True
        
    row_fill = ADOPT_ROW_FILL if is_adopt else MAINTAIN_ROW_FILL
    status_font = STATUS_ADOPT_FONT if is_adopt else STATUS_MAINTAIN_FONT
    
    ws2.row_dimensions[r].height = 24
    for c_idx in range(1, 11):
        cell = ws2.cell(row=r, column=c_idx)
        cell.fill = row_fill
        cell.font = REGULAR_FONT
        cell.border = THIN_BORDER
        if c_idx in [1, 5]: cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        elif c_idx == 6:
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.font = status_font
        elif c_idx in [7, 8]: cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        else: cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

wb.save(file_path)
print("✅ PolyU feedback sheet styling fully updated and saved!")

desktop_bkp = r"c:\Users\rtx\Desktop\香港理工大学_数据核对与差异反馈表.xlsx"
try:
    wb.save(desktop_bkp)
    print("✅ Synced to Desktop root:", desktop_bkp)
except Exception as e:
    print("Notice on desktop root:", e)
